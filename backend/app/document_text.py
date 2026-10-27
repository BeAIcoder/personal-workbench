"""聊天文档附件的有界正文提取与旧版 Office 转换。"""
import os
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

DOCUMENT_SUFFIXES = frozenset({".pdf", ".docx", ".doc", ".pptx", ".ppt"})
MAX_DOCUMENT_TEXT_CHARS = 80_000
MAX_PDF_PAGES = 300
MAX_PPT_SLIDES = 500
MAX_OOXML_MEMBERS = 5_000
MAX_OOXML_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
LEGACY_CONVERSION_TIMEOUT_SECONDS = 60


class DocumentExtractionError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractionResult:
    characters: int
    truncated: bool
    converted: bool


def extracted_text_path(source: Path) -> Path:
    source = Path(source)
    return source.parent / ".extracted" / f"{source.name}.txt"


class _TextAccumulator:
    def __init__(self):
        self.limit = MAX_DOCUMENT_TEXT_CHARS
        self.parts: list[str] = []
        self.size = 0
        self.truncated = False

    def add(self, value: str) -> None:
        if self.truncated:
            return
        block = str(value or "").strip()
        if not block:
            return
        separator_size = 2 if self.parts else 0
        remaining = self.limit - self.size - separator_size
        if remaining <= 0:
            self.truncated = True
            return
        if len(block) > remaining:
            block = block[:remaining]
            self.truncated = True
        if self.parts:
            self.parts.append("\n\n")
            self.size += 2
        self.parts.append(block)
        self.size += len(block)

    def finish(self) -> str:
        if not self.parts:
            return ""
        text = "".join(self.parts)
        if not self.truncated:
            return text
        marker = "[内容过长，已截断]"
        separator = "\n\n"
        keep = max(0, self.limit - len(marker) - len(separator))
        text = text[:keep].rstrip()
        result = f"{text}{separator}{marker}"
        return result[:self.limit]


def _extract_pdf(source: Path) -> _TextAccumulator:
    import pdfplumber

    accumulator = _TextAccumulator()
    with pdfplumber.open(str(source)) as pdf:
        for index, page in enumerate(pdf.pages):
            if index >= MAX_PDF_PAGES:
                accumulator.truncated = True
                break
            text = (page.extract_text() or "").strip()
            if text:
                accumulator.add(f"[第 {index + 1} 页]\n{text}")
            if accumulator.truncated:
                break
    return accumulator


def _iter_docx_blocks(document):
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, document).text.strip()
            if text:
                yield text
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = ["\n".join(p.text for p in cell.paragraphs).strip() for cell in row.cells]
                line = "\t".join(cells).strip()
                if line:
                    yield line


def _extract_docx(source: Path) -> _TextAccumulator:
    from docx import Document

    accumulator = _TextAccumulator()
    document = Document(str(source))
    for block in _iter_docx_blocks(document):
        accumulator.add(block)
        if accumulator.truncated:
            break
    return accumulator


def _shape_blocks(shape) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        blocks = []
        for child in shape.shapes:
            blocks.extend(_shape_blocks(child))
        return blocks
    blocks = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        if text:
            blocks.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            line = "\t".join(cell.text.strip() for cell in row.cells).strip()
            if line:
                blocks.append(line)
    return blocks


def _extract_pptx(source: Path) -> _TextAccumulator:
    from pptx import Presentation

    accumulator = _TextAccumulator()
    presentation = Presentation(str(source))
    for index, slide in enumerate(presentation.slides):
        if index >= MAX_PPT_SLIDES:
            accumulator.truncated = True
            break
        blocks = []
        for shape in slide.shapes:
            blocks.extend(_shape_blocks(shape))
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ""
            if notes:
                blocks.append(f"[备注]\n{notes}")
        if blocks:
            accumulator.add(f"[幻灯片 {index + 1}]\n" + "\n".join(blocks))
        if accumulator.truncated:
            break
    return accumulator


def _validate_ooxml(source: Path) -> None:
    try:
        with zipfile.ZipFile(source) as archive:
            members = archive.infolist()
    except (OSError, zipfile.BadZipFile) as exc:
        raise DocumentExtractionError("无法解析 Office 文档：文件不是有效的 OOXML 格式") from exc
    total_size = sum(member.file_size for member in members)
    if len(members) > MAX_OOXML_MEMBERS or total_size > MAX_OOXML_UNCOMPRESSED_BYTES:
        raise DocumentExtractionError("Office 文档解压后超过安全上限，请拆分或压缩后再上传")


def _find_soffice() -> str | None:
    candidates = [
        os.environ.get("SOFFICE_PATH"),
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice",
        "/usr/bin/libreoffice",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(Path(candidate))
    return None


def _convert_legacy_to_ooxml(source: Path, work_dir: Path) -> Path:
    suffix = source.suffix.lower()
    target_suffix = {".doc": ".docx", ".ppt": ".pptx"}.get(suffix)
    if target_suffix is None:
        raise DocumentExtractionError("旧版转换仅支持 .doc 与 .ppt")
    executable = _find_soffice()
    if executable is None:
        raise DocumentExtractionError("旧版 .doc/.ppt 需要 LibreOffice 转换；请安装 LibreOffice 或另存为新版格式")

    work_dir = Path(work_dir)
    output_dir = work_dir / "output"
    profile_dir = work_dir / "profile"
    output_dir.mkdir(parents=True, exist_ok=True)
    profile_dir.mkdir(parents=True, exist_ok=True)
    command = [
        executable,
        "--headless",
        "--nologo",
        "--nodefault",
        "--nofirststartwizard",
        "--nolockcheck",
        f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
        "--convert-to",
        target_suffix.lstrip("."),
        "--outdir",
        str(output_dir.resolve()),
        str(source.resolve()),
    ]
    try:
        completed = subprocess.run(
            command,
            stdin=subprocess.DEVNULL,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=LEGACY_CONVERSION_TIMEOUT_SECONDS,
            check=False,
            shell=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except subprocess.TimeoutExpired as exc:
        raise DocumentExtractionError("LibreOffice 转换超时，请拆分文档或另存为新版格式后重试") from exc
    except OSError as exc:
        raise DocumentExtractionError("无法启动 LibreOffice 转换器") from exc
    if completed.returncode != 0:
        raise DocumentExtractionError("LibreOffice 转换失败，请另存为 .docx/.pptx 后重试")

    output = output_dir / f"{source.stem}{target_suffix}"
    try:
        output.resolve().relative_to(output_dir.resolve())
    except ValueError as exc:
        raise DocumentExtractionError("LibreOffice 输出路径越界") from exc
    if not output.is_file():
        raise DocumentExtractionError("LibreOffice 未生成预期的新版文档，请另存为新版格式后重试")
    return output


def _write_text_atomic(destination: Path, text: str) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=f"{destination.name}.", suffix=".tmp", dir=destination.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(text)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)


def extract_document(source: Path, destination: Path) -> ExtractionResult:
    source = Path(source)
    destination = Path(destination)
    suffix = source.suffix.lower()
    try:
        if suffix not in DOCUMENT_SUFFIXES:
            raise DocumentExtractionError(f"不支持的文档类型：{suffix or '无扩展名'}")
        converted = False
        with tempfile.TemporaryDirectory(prefix="workbench-document-") as temporary_dir:
            actual_source = source
            if suffix in {".doc", ".ppt"}:
                actual_source = _convert_legacy_to_ooxml(source, Path(temporary_dir))
                converted = True
            actual_suffix = actual_source.suffix.lower()
            if actual_suffix in {".docx", ".pptx"}:
                _validate_ooxml(actual_source)
            if actual_suffix == ".pdf":
                accumulator = _extract_pdf(actual_source)
            elif actual_suffix == ".docx":
                accumulator = _extract_docx(actual_source)
            elif actual_suffix == ".pptx":
                accumulator = _extract_pptx(actual_source)
            else:
                raise DocumentExtractionError("转换后的文档格式不受支持")
            text = accumulator.finish()
            if not text.strip():
                raise DocumentExtractionError("文档未提取到文字；扫描件或纯图片文档请先进行 OCR")
            _write_text_atomic(destination, text)
        return ExtractionResult(characters=len(text), truncated=accumulator.truncated, converted=converted)
    except DocumentExtractionError:
        destination.unlink(missing_ok=True)
        raise
    except Exception as exc:
        destination.unlink(missing_ok=True)
        raise DocumentExtractionError(f"无法解析 {suffix} 文档，请确认文件格式正确且未损坏") from exc
