"""finance_expense 模块业务逻辑：费用分类引擎定位、分类执行、运行记录与 Agent 工具。

引擎复用（L0 直挂）：真实引擎在 01-财税自动化/现金流滚动预测/费用分类汇总系统-v2.0/
（expense_classification.py 的 ExpenseClassifier + rules_loader.py + classification_rules.json）。
规则 JSON 是单一权威源，本模块只读引用，不复制其内容。

故障隔离：引擎目录缺失或导入失败时，模块本身照常加载（manifest/路由/表均不受影响），
依赖引擎的接口返回 503 + 中文提示；运行记录查询等接口不受引擎状态影响。

行数护栏：上传 Excel 先按共享上限 MAX_IMPORT_ROWS 统计各工作表非空行，超限在引擎执行
与运行记录写库前失败关闭（400），超大文件不进入引擎解析。
"""
import contextlib
from contextlib import contextmanager
import importlib
import os
import sys
import threading
import time
from pathlib import Path
from typing import Any

from sqlalchemy import func
from sqlalchemy.orm import Session

from ...config import BASE_DIR, get_finance_data_dir, settings
from ...database import SessionLocal
from ...utils import resolve_engine_root
from .models import ExpenseRow, ExpenseRun

# 单次导入行数上限由 core 任务提供（app/import_limits.py）；缺失时失败关闭，
# 绝不放行无上限文件进入引擎（与 BIP/税务/数据仓等财务模块同一口径）。
try:
    from ...import_limits import MAX_IMPORT_ROWS
except ImportError:  # pragma: no cover - 集成共享模块后不存在该分支
    MAX_IMPORT_ROWS = None

MODULE_NAME = "finance_expense"

# 引擎默认相对位置：01-财税自动化/现金流滚动预测/费用分类汇总系统-v2.0
_ENGINE_RELATIVE = Path("现金流滚动预测") / "费用分类汇总系统-v2.0"
_ENGINE_REQUIRED_FILES = ("expense_classification.py", "rules_loader.py", "classification_rules.json")

# 上传约束：扩展名白名单 + 大小上限（测试可 monkeypatch 调低）
ALLOWED_EXTENSIONS = {".xlsx", ".xls"}
MAX_UPLOAD_BYTES = 20 * 1024 * 1024

# 银行明细账读取失败提示：行数护栏的 Excel 统计与引擎 read_bank_data 返回 None 共用同一口径
_BANK_READ_FAIL_DETAIL = (
    "银行明细账读取失败：请确认上传的是银行存款辅助明细账"
    "（需含 科目编码/日期/摘要/贷方本币 等列；表头行已自动定位）"
)


class EngineUnavailableError(RuntimeError):
    """费用分类引擎不可用（目录缺失或导入失败），由路由层映射为 503。"""


# 引擎加载缓存：仅缓存成功结果；失败不缓存，下次调用重新探测（目录后补上也无需重启）
_ENGINE: dict[str, Any] = {"classifier_cls": None, "load_rules": None, "app_config": None, "engine_dir": ""}
_BANK_HEADER_LOCK = threading.Lock()
# 引擎加载锁：os.chdir 是进程级状态，并发请求下 chdir+import 临界区必须串行化
_ENGINE_LOCK = threading.Lock()


def _reset_engine_state() -> None:
    """清空引擎缓存（测试隔离与故障恢复用）。"""
    _ENGINE.update(classifier_cls=None, load_rules=None, app_config=None, engine_dir="")


def _row_has_value(values) -> bool:
    """一行是否含非空值：None/空白字符串视为空（不同读取器的空单元格表示不统一，这里归一）。"""
    return any(v is not None and str(v).strip() != "" for v in values)


def _count_xlsx_rows(path: Path, limit: int) -> int:
    """.xlsx 非空行计数：openpyxl 只读模式流式迭代，超过 limit 即停，不完整解析无限行。

    reset_dimensions 防尺寸标签失真：只读迭代默认被工作表 <dimension> 声明行数截断
    （pandas 内部同样处理），标签失真的文件会漏数后续数据行。
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        total = 0
        for sheet in workbook.worksheets:
            sheet.reset_dimensions()
            for row in sheet.iter_rows(values_only=True):
                if _row_has_value(row):
                    total += 1
                    if total > limit:
                        return total
        return total
    finally:
        workbook.close()


def _count_xls_rows(path: Path, limit: int) -> int:
    """.xls 非空行计数：xlrd 无行流式读取，按工作表逐张解析（on_demand 及时释放），
    超过 limit 即停止解析剩余工作表；非空行口径与 .xlsx 一致。"""
    import xlrd

    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        total = 0
        for idx in range(book.nsheets):
            sheet = book.sheet_by_index(idx)
            for r in range(sheet.nrows):
                if _row_has_value(sheet.cell_value(r, c) for c in range(sheet.ncols)):
                    total += 1
                    if total > limit:
                        return total
            book.unload_sheet(idx)
        return total
    finally:
        book.release_resources()


def _excel_row_count(path: Path, limit: int) -> int:
    """Excel 各工作表非空行合计（含表头行、跳过标题空区）：未超限时精确，超限早停于 limit+1。"""
    if path.suffix.lower() == ".xls":
        return _count_xls_rows(path, limit)
    return _count_xlsx_rows(path, limit)


def _guard_import_rows(path: Path) -> None:
    """行数护栏（复用共享 MAX_IMPORT_ROWS）：超限/上限服务缺失即失败关闭，抛 ValueError。

    在 classify_file 最前端调用：先于引擎执行与运行记录写库，超限文件不产生任何半成品；
    Excel 读取失败与引擎 read_bank_data 返回 None 同口径（提示确认银行明细账格式）。
    """
    if MAX_IMPORT_ROWS is None:
        raise ValueError("导入行数上限服务暂不可用，请稍后重试")
    try:
        count = _excel_row_count(path, MAX_IMPORT_ROWS)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(_BANK_READ_FAIL_DETAIL) from exc
    if count > MAX_IMPORT_ROWS:
        raise ValueError(
            f"文件数据行数（{count}）超过单次导入上限（{MAX_IMPORT_ROWS} 行），请拆分后导入"
        )


def module_subdir(kind: str) -> Path:
    """模块数据子目录：get_finance_data_dir()/finance_expense/inbox|outbox|tmp。"""
    if kind not in ("inbox", "outbox", "tmp"):
        raise ValueError(f"未知模块子目录类型：{kind}")
    path = get_finance_data_dir() / MODULE_NAME / kind
    path.mkdir(parents=True, exist_ok=True)
    return path


def _is_engine_dir(path: Path) -> bool:
    """目录须同时具备引擎入口、规则加载器与规则 JSON 权威源。"""
    return path.is_dir() and all((path / name).is_file() for name in _ENGINE_REQUIRED_FILES)


def find_engine_dir() -> Path | None:
    """定位费用分类引擎目录。

    settings.finance_expense_engine_path 非空时只用配置值（无效即不可用，不回退）；
    否则从 BASE_DIR 向上找「01-财税自动化」目录，推导其下的 现金流滚动预测/费用分类汇总系统-v2.0。
    """
    configured = (settings.finance_expense_engine_path or "").strip()
    if configured:
        # 支持目录或单一文件（文件取其父目录）——见 utils.resolve_engine_root
        candidate = resolve_engine_root(configured)
        return candidate if candidate and _is_engine_dir(candidate) else None
    for parent in BASE_DIR.parents:
        if parent.name == "01-财税自动化":
            candidate = parent / _ENGINE_RELATIVE
            return candidate if _is_engine_dir(candidate) else None
    return None


def _locate_bank_header_row(path: Path) -> int | None:
    """定位银行明细账表头行（0-based）：前 25 行内找含「科目编码」且带「日期」「摘要」的行。

    用友导出的元数据块行数不固定（本次实测 12 行），引擎的 bank_header_row 是固定
    配置——由模块侧动态定位后覆盖，兼容不同导出版本（外部验证第 4 章 E1）。
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=25, values_only=True)):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if "科目编码" in cells and any("日期" in c for c in cells) and any("摘要" in c for c in cells):
                    return i
        return None
    finally:
        wb.close()


@contextmanager
def _bank_header_override(engine: dict, row: int | None):
    """read_bank_data 期间把引擎配置的 bank_header_row 覆盖为定位值（用后恢复）。

    引擎 _cfg 在调用时读其配置（app_config 可能是 dict，也可能是引擎的 config
    模块——FILE_FORMAT_CONFIG 为其中的 dict，直接改其键值即可）；并发分类请求
    会争用同一份配置，故整段（覆盖+调用）持模块锁串行，恢复保证在 finally。
    """
    cfg = engine.get("app_config")
    if isinstance(cfg, dict):
        ff = cfg.get("FILE_FORMAT_CONFIG")
    else:
        ff = getattr(cfg, "FILE_FORMAT_CONFIG", None)
    if row is None or not isinstance(ff, dict):
        yield
        return
    original = ff.get("bank_header_row")
    with _BANK_HEADER_LOCK:
        ff["bank_header_row"] = row
        try:
            yield
        finally:
            ff["bank_header_row"] = original


def _ensure_engine() -> dict[str, Any]:
    """懒加载引擎：首次调用时定位目录、插入 sys.path 并 import，成功后缓存。

    引擎 import 的副作用处理：
    - expense_classification 顶层按相对路径创建 expense_classification.log，
      import 期间临时把 CWD 切到模块 tmp 目录，避免污染宿主工作目录；
    - 引擎以顶层 `import config` 读取自带 config.py，import 期间临时遮蔽
      sys.modules 里可能存在的同名外来模块，import 完成后恢复原状（引擎内部
      已绑定自身 config 引用，本模块加载规则时也显式传入该引用）。

    并发安全：模块级 _ENGINE_LOCK + 双重检查，把 chdir+import+恢复 整个临界区
    串行化（CWD 是进程级状态，并发 chdir 会互相漂移）。
    """
    if _ENGINE["classifier_cls"] is not None:
        return _ENGINE

    with _ENGINE_LOCK:
        if _ENGINE["classifier_cls"] is not None:  # 双重检查：等锁期间可能已被其他线程加载
            return _ENGINE

        engine_dir = find_engine_dir()
        if engine_dir is None:
            raise EngineUnavailableError(
                "费用分类引擎不可用：未找到引擎目录（现金流滚动预测/费用分类汇总系统-v2.0），"
                "请检查 FINANCE_EXPENSE_ENGINE_PATH 配置或确认引擎项目在位"
            )

        tmp_dir = module_subdir("tmp")
        engine_key = str(engine_dir)
        prev_cwd = os.getcwd()
        stashed_config = sys.modules.pop("config", None)
        try:
            if engine_key in sys.path:
                sys.path.remove(engine_key)
            sys.path.insert(0, engine_key)
            os.chdir(tmp_dir)
            try:
                ec_mod = importlib.import_module("expense_classification")
                rl_mod = importlib.import_module("rules_loader")
            finally:
                os.chdir(prev_cwd)
        except EngineUnavailableError:
            raise
        except Exception as exc:
            raise EngineUnavailableError(f"费用分类引擎加载失败：{exc}") from exc
        finally:
            # 恢复 sys.modules 中的外来 config（若有）；引擎自带的 config 引用已绑定在模块内
            sys.modules.pop("config", None)
            if stashed_config is not None:
                sys.modules["config"] = stashed_config

        _ENGINE.update(
            classifier_cls=ec_mod.ExpenseClassifier,
            load_rules=rl_mod.load_classification_rules,
            app_config=getattr(ec_mod, "_app_config", None),
            engine_dir=engine_key,
        )
        return _ENGINE


def ensure_engine_ready() -> None:
    """供路由层快速失败的引擎可用性探测（不可用抛 EngineUnavailableError）。"""
    _ensure_engine()


def _load_rules() -> dict:
    """按引擎同口径加载分类规则（JSON 权威源 + 引擎自带 config 覆盖），每次调用重读文件。"""
    engine = _ensure_engine()
    return engine["load_rules"](config_module=engine["app_config"])


def run_to_dict(run: ExpenseRun) -> dict:
    """运行记录转响应 dict。"""
    return {
        "id": run.id,
        "original_filename": run.original_filename,
        "total_rows": run.total_rows,
        "matched_rows": run.matched_rows,
        "status": run.status,
        "duration_sec": run.duration_sec,
        "artifact_path": run.artifact_path,
        "created_at": run.created_at,
    }


def classify_file(db: Session, original_filename: str, stored_path: Path) -> ExpenseRun:
    """执行一轮费用分类：读取银行明细账 → 整合 → 清洗 → 分类 → 汇总 → 导出，结果入库。

    行数护栏先行：超过共享 MAX_IMPORT_ROWS 的文件在统计阶段即失败关闭（ValueError，
    路由层映射 400 并清理 inbox 上传），引擎不执行、不写运行记录、无半成品。
    失败时运行记录落库为 failed 并把中文错误抛给路由层（400）。
    引擎不可用抛 EngineUnavailableError（503），此时不写运行记录。
    """
    _guard_import_rows(stored_path)
    engine = _ensure_engine()
    run = ExpenseRun(original_filename=original_filename[:300], status="running")
    db.add(run)
    db.commit()
    db.refresh(run)

    started = time.perf_counter()
    try:
        clf = engine["classifier_cls"]()
        header_row = _locate_bank_header_row(stored_path)
        with _bank_header_override(engine, header_row):
            if clf.read_bank_data(str(stored_path)) is None:
                raise ValueError(_BANK_READ_FAIL_DETAIL)
        if clf.integrate_data() is None:
            raise ValueError("数据整合失败")
        if clf.clean_data() is None:
            raise ValueError("数据清洗失败")
        classified = clf.classify_expenses()
        if classified is None:
            raise ValueError("费用分类执行失败")
        if clf.summarize_expenses() is None:
            raise ValueError("费用汇总计算失败")
        out_name = f"run_{run.id}_费用分类结果.xlsx"
        out_path = module_subdir("outbox") / out_name
        if clf.export_results(str(out_path)) is None:
            raise ValueError("结果报表导出失败")

        rows: list[ExpenseRow] = []
        matched = 0
        for row_no, (_, record) in enumerate(classified.iterrows(), start=1):
            category = str(record.get("费用类别") or "未分类")
            hit = category != "未分类"
            matched += int(hit)
            rows.append(
                ExpenseRow(
                    run_id=run.id,
                    row_no=row_no,
                    summary=str(record.get("摘要") or "")[:500],
                    amount=float(record.get("贷方本币") or 0),
                    category=category[:50],
                    hit=hit,
                )
            )
        db.add_all(rows)
        run.status = "success"
        run.total_rows = len(rows)
        run.matched_rows = matched
        run.duration_sec = round(time.perf_counter() - started, 3)
        run.artifact_path = f"{MODULE_NAME}/outbox/{out_name}"
        db.commit()
        db.refresh(run)
        return run
    except Exception:
        run.status = "failed"
        run.duration_sec = round(time.perf_counter() - started, 3)
        db.commit()
        raise


def list_runs(db: Session, page: int, page_size: int) -> dict:
    """分页查询运行记录（按 id 倒序）。"""
    total = db.query(func.count(ExpenseRun.id)).scalar() or 0
    items = (
        db.query(ExpenseRun)
        .order_by(ExpenseRun.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    return {"total": total, "items": [run_to_dict(r) for r in items]}


def get_run(db: Session, run_id: int) -> ExpenseRun | None:
    """按 id 取运行记录。"""
    return db.get(ExpenseRun, run_id)


def run_detail(db: Session, run: ExpenseRun) -> dict:
    """运行详情：记录本体 + 按命中分类聚合的金额/行数统计 + 全部明细行。"""
    rows = db.query(ExpenseRow).filter(ExpenseRow.run_id == run.id).order_by(ExpenseRow.row_no).all()
    stats_map: dict[str, list] = {}
    for r in rows:
        bucket = stats_map.setdefault(r.category, [0, 0.0])
        bucket[0] += 1
        bucket[1] += r.amount
    stats = [
        {"category": category, "count": v[0], "amount": round(v[1], 2)}
        for category, v in sorted(stats_map.items(), key=lambda kv: kv[1][1], reverse=True)
    ]
    return {
        "run": run_to_dict(run),
        "stats": stats,
        "rows": [
            {
                "id": r.id,
                "row_no": r.row_no,
                "summary": r.summary,
                "amount": r.amount,
                "category": r.category,
                "hit": r.hit,
            }
            for r in rows
        ],
    }


def resolve_artifact(run: ExpenseRun) -> Path | None:
    """解析运行产物路径：必须落在模块 outbox 内（防路径穿越），文件须存在。"""
    rel = (run.artifact_path or "").strip()
    if not rel:
        return None
    base = get_finance_data_dir().resolve()
    candidate = (base / rel).resolve()
    outbox = (base / MODULE_NAME / "outbox").resolve()
    if outbox not in candidate.parents:
        return None
    return candidate if candidate.is_file() else None


def rules_summary() -> dict:
    """规则摘要：分类条数 / 关键词总数 / 来源路径（引擎不可用时抛 503）。"""
    engine = _ensure_engine()
    rules = _load_rules()
    return {
        "engine_dir": engine["engine_dir"],
        "rules_file": str(Path(engine["engine_dir"]) / "classification_rules.json"),
        "total_categories": len(rules),
        "total_keywords": sum(len(rule.get("keywords", [])) for rule in rules.values()),
        "categories": list(rules.keys()),
    }


def reload_rules() -> dict:
    """重载规则：rules_loader 每次调用都重读 JSON 文件，重新加载并返回最新摘要即可。"""
    return {"ok": True, "summary": rules_summary()}


def expense_runs_summary(limit: int = 5) -> dict:
    """Agent 工具（readonly）：最近 N 次费用分类运行统计。"""
    limit = max(1, min(int(limit), 20))
    with SessionLocal() as db:
        total = db.query(func.count(ExpenseRun.id)).scalar() or 0
        runs = db.query(ExpenseRun).order_by(ExpenseRun.id.desc()).limit(limit).all()
        recent = [
            {
                "id": r.id,
                "original_filename": r.original_filename,
                "total_rows": r.total_rows,
                "matched_rows": r.matched_rows,
                "命中率": f"{(r.matched_rows / r.total_rows * 100):.1f}%" if r.total_rows else "-",
                "status": r.status,
                "duration_sec": r.duration_sec,
                "created_at": r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else "",
            }
            for r in runs
        ]
    return {"ok": True, "total_runs": total, "recent": recent}
