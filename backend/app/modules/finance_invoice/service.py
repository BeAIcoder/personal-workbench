"""finance_invoice 模块业务逻辑与 Agent 工具函数。

职责：
- 上传批次落盘与解析编排：PDF 文本提取 / 字段正则解析直挂发票工具引擎（engine.py）；
- 销方名称补充提取：引擎只解析购方，销方由本模块正则补齐；
- 去重（发票号码+校验码幂等，批内 + 跨批次）与冲红配对（同购方、金额绝对值相等、一对一），
  规则与引擎 FileHandler.deduplicate_invoices / match_red_invoices 保持一致；
- 导出：按引擎重命名规则（{购方}{品牌}{金额}{类型}.pdf、{年月}{项目} 目录）复制 PDF、
  生成每项目「发票清单.xlsx」与「汇总清单.xlsx」，并打包 ZIP；
- 数据目录：get_finance_data_dir()/finance_invoice/{inbox,outbox,tmp}；
- Agent 工具只开放查询类（invoice_batch_stats），出错返回 {"ok": False, "error": ...}。
"""
from __future__ import annotations

import io
import json
import logging
import re
import shutil
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func
from sqlalchemy.orm import Session

from ...config import get_finance_data_dir
from ...database import SessionLocal
from . import engine as engine_mod
from .models import InvoiceBatch, InvoiceItem

logger = logging.getLogger(__name__)

# ---- 上传约束（扩展名白名单 + 大小限制，与宿主上传口径一致）----
ALLOWED_SUFFIXES = {".pdf", ".zip"}
MAX_FILE_BYTES = 20 * 1024 * 1024  # 单文件 20MB
MAX_BATCH_BYTES = 100 * 1024 * 1024  # 单批次合计 100MB（ZIP 按解压后口径）
MAX_FILES_PER_BATCH = 200

# 销方名称补充提取（引擎仅解析购方）：匹配清洁后文本中的「销 名称：xxx」等写法
_SELLER_PATTERNS = [
    re.compile(r"销\s+名称[：:]\s*([^\s]{2,30})"),
    re.compile(r"销\s*售\s*方[^\n\r]{0,20}?名称[：:]\s*([^\s]{2,30})"),
    re.compile(r"销售方[^\n\r]{0,20}?名称[：:]\s*([^\s]{2,30})"),
    re.compile(r"卖\s*方[^\n\r]{0,20}?名称[：:]\s*([^\s]{2,30})"),
]
_SELLER_STOP_WORDS = ("发票", "代码", "号码", "校验", "金额", "合计", "人民币", "税率", "税额")

# 金额 / 日期 / 文件名的处理规则镜像引擎 invoice_processor/src/utils.py（保持两处口径一致）
_YM_PATTERNS = [
    re.compile(r"(\d{4})年(\d{1,2})月\d{1,2}日"),
    re.compile(r"(\d{4})-(\d{1,2})-\d{1,2}"),
    re.compile(r"(\d{4})\.(\d{1,2})\.\d{1,2}"),
    re.compile(r"(\d{4})/(\d{1,2})/\d{1,2}"),
]
_ILLEGAL_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1F]')
# 同引擎口径：保留中英文、数字、连字符与下划线（一-鿿 即 CJK 统一表意区间 4E00-9FFF）
_NON_WORD_CHARS = re.compile(r"[^\w一-鿿\-_]")
_AMOUNT_NUMBER_FORMAT = "#,##0.00"

# 可参与导出 / 冲红配对的明细状态
_ACTIVE_ITEM_STATUS = (InvoiceItem.STATUS_PARSED, InvoiceItem.STATUS_INCOMPLETE)


@dataclass
class ParsedInvoice:
    """单张发票的解析结果（引擎 InvoiceInfo 的模块内表示 + 销方名称）。"""

    invoice_number: str = ""
    check_code: str = ""
    buyer_name: str = ""
    seller_name: str = ""
    amount: str = ""
    tax_amount: str = ""
    tax_rate: str = ""
    invoice_date: str = ""
    invoice_type: str = ""
    invoice_direction: str = ""
    project_name: str = ""
    brand_name: str = ""
    remark: str = ""
    extraction_success: bool = True


# ---------------- 数据目录 ----------------

def module_data_dir() -> Path:
    """模块数据目录：get_finance_data_dir()/finance_invoice。"""
    return get_finance_data_dir() / "finance_invoice"


def _ensure_dirs() -> dict[str, Path]:
    base = module_data_dir()
    dirs = {"inbox": base / "inbox", "outbox": base / "outbox", "tmp": base / "tmp"}
    for d in dirs.values():
        d.mkdir(parents=True, exist_ok=True)
    return dirs


# ---------------- 纯函数工具（镜像引擎口径） ----------------

def parse_amount(amount_str: Optional[str]) -> float:
    """金额原文 → float；支持千分位逗号、货币符号、负数（冲红），失败返回 0.0。"""
    if not amount_str:
        return 0.0
    try:
        cleaned = str(amount_str).replace(",", "").replace("¥", "").replace("￥", "")
        cleaned = cleaned.replace(" ", "").strip()
        if not cleaned:
            return 0.0
        return float(cleaned)
    except (ValueError, TypeError):
        return 0.0


def sanitize_filename(name: str) -> str:
    """剔除文件名非法字符（同引擎 sanitize_filename）。"""
    name = _ILLEGAL_FILENAME_CHARS.sub("", name or "").strip()
    return name or "未命名"


def parse_date_to_year_month(date_str: str) -> str:
    """开票日期 → 「yy年M月」（同引擎 parse_date_to_year_month，用于导出目录命名）。"""
    if not date_str:
        return ""
    for pattern in _YM_PATTERNS:
        match = pattern.search(date_str)
        if match:
            year, month = match.group(1), match.group(2)
            short_year = year[-2:] if len(year) == 4 else year
            return f"{short_year}年{int(month)}月"
    return ""


def extract_seller_name(text: str) -> str:
    """从清洁后的发票文本补充提取销方名称（引擎只解析购方）；失败返回空串。"""
    for pattern in _SELLER_PATTERNS:
        match = pattern.search(text)
        if not match:
            continue
        name = re.sub(r"\s+", "", match.group(1)).strip(" :：")
        if 1 < len(name) <= 30 and not any(kw in name for kw in _SELLER_STOP_WORDS):
            return name
    return ""


def dedup_key(invoice_number: str, check_code: str) -> str:
    """去重幂等键：发票号码 + 校验码（同引擎 FileHandler.deduplicate_invoices）；无有效号码返回空串。"""
    number = (invoice_number or "").strip()
    if not number or number == "未知":
        return ""
    code = (check_code or "").strip()
    return f"{number}|{code}" if code else f"{number}|"


def build_renamed_filename(item: InvoiceItem) -> str:
    """引擎 FileHandler.rename_invoice_file 同款命名：{购方}{品牌}{金额}{类型}.pdf，超长截断。"""
    buyer = _NON_WORD_CHARS.sub("", sanitize_filename((item.buyer_name or "")[:15]))
    brand = _NON_WORD_CHARS.sub("", sanitize_filename(item.brand_name or ""))[:10] if item.brand_name else ""
    amount = re.sub(r"[^\d.]", "", (item.amount or "").replace(",", "")) or "0"
    inv_type = item.invoice_type or ""
    if brand:
        name = f"{buyer}{brand}{amount}{inv_type}.pdf"
    else:
        name = f"{buyer}{amount}{inv_type}.pdf"
    if len(name) > 100:
        if brand:
            name = f"{buyer[:8]}{brand[:8]}{amount}{inv_type}.pdf"
        else:
            name = f"{buyer[:10]}{amount}{inv_type}.pdf"
    return name


# ---------------- 引擎解析 ----------------

def _require_engine() -> None:
    """批次创建前置检查：引擎不可用抛 EngineUnavailableError（路由层转 503）。"""
    engine_mod.get_engine()


def parse_invoice_text(processor, text: str) -> ParsedInvoice:
    """对已提取的 PDF 文本做字段解析：引擎正则 + 本模块销方补充。processor 为引擎 PDFInvoiceProcessor。"""
    cleaned = processor._clean_text(text)
    info = processor._parse_invoice_info(cleaned)
    return ParsedInvoice(
        invoice_number=(info.invoice_number or "")[:50],
        check_code=(info.check_code or "")[:50],
        buyer_name=(info.buyer_name or "")[:200],
        seller_name=extract_seller_name(cleaned)[:200],
        amount=(info.amount or "")[:50],
        tax_amount=(info.tax_amount or "")[:50],
        tax_rate=(info.tax_rate or "")[:20],
        invoice_date=(info.invoice_date or "")[:30],
        invoice_type=(info.invoice_type or "")[:10],
        invoice_direction=(info.invoice_direction or "")[:10],
        project_name=(info.project_name or "")[:100],
        brand_name=(info.brand_name or "")[:50],
        remark=(info.remarks or "")[:500],
        extraction_success=bool(info.extraction_success),
    )


def _parse_pdf_with_engine(pdf_path: Path) -> Optional[ParsedInvoice]:
    """用引擎解析单个 PDF；解析失败返回 None（不抛异常，单文件失败不拖垮批次）。"""
    handle = engine_mod.get_engine()
    processor = handle.processor_cls(handle.config)
    try:
        text = processor._extract_text_once(str(pdf_path))
    except Exception as exc:
        logger.warning("PDF 文本提取失败 %s：%s", pdf_path, exc)
        return None
    if not text.strip():
        logger.warning("PDF 文本为空：%s", pdf_path)
        return None
    cleaned = processor._clean_text(text)
    if not cleaned.strip():
        return None
    try:
        info = processor._parse_invoice_info(cleaned)
    except Exception as exc:
        logger.warning("发票字段解析失败 %s：%s", pdf_path, exc)
        return None
    return ParsedInvoice(
        invoice_number=(info.invoice_number or "")[:50],
        check_code=(info.check_code or "")[:50],
        buyer_name=(info.buyer_name or "")[:200],
        seller_name=extract_seller_name(cleaned)[:200],
        amount=(info.amount or "")[:50],
        tax_amount=(info.tax_amount or "")[:50],
        tax_rate=(info.tax_rate or "")[:20],
        invoice_date=(info.invoice_date or "")[:30],
        invoice_type=(info.invoice_type or "")[:10],
        invoice_direction=(info.invoice_direction or "")[:10],
        project_name=(info.project_name or "")[:100],
        brand_name=(info.brand_name or "")[:50],
        remark=(info.remarks or "")[:500],
        extraction_success=bool(info.extraction_success),
    )


# ---------------- 去重与冲红配对 ----------------

def existing_dedup_keys(db: Session, exclude_batch_id: Optional[int] = None) -> set[str]:
    """库内既有有效发票的幂等键集合（跨批次去重用）。"""
    query = db.query(InvoiceItem.invoice_number, InvoiceItem.check_code).filter(
        InvoiceItem.status.in_(_ACTIVE_ITEM_STATUS)
    )
    if exclude_batch_id is not None:
        query = query.filter(InvoiceItem.batch_id != exclude_batch_id)
    return {k for k in (dedup_key(n, c) for n, c in query.all()) if k}


def mark_duplicates(items: list[InvoiceItem], existing_keys: Optional[set[str]] = None) -> int:
    """批内 + 跨批次去重：后来者标记为 duplicate，返回重复个数。

    规则同引擎：号码+校验码构成幂等键，保留首次出现；failed 明细不参与。
    """
    seen = set(existing_keys or set())
    duplicates = 0
    for item in items:
        if item.status == InvoiceItem.STATUS_FAILED:
            continue
        key = dedup_key(item.invoice_number, item.check_code)
        if not key:
            continue
        if key in seen:
            item.status = InvoiceItem.STATUS_DUPLICATE
            duplicates += 1
        else:
            seen.add(key)
    return duplicates


def pair_red_invoices(items: list[InvoiceItem]) -> tuple[int, int]:
    """冲红配对：同购方 + 金额绝对值相等 + 一对一（同引擎 match_red_invoices）。

    配对成功双向写入 paired_item_id；返回 (配对成功的对数, 未配对冲红张数)。
    仅 parsed/incomplete 明细参与（重复与失败明细不配对）。
    """
    candidates = [i for i in items if i.status in _ACTIVE_ITEM_STATUS]
    reds = [i for i in candidates if i.invoice_direction == "冲红发票"]
    positives = [i for i in candidates if i.invoice_direction == "正数发票"]
    used_positive: set[int] = set()
    matched = 0
    for red in reds:
        red_abs = abs(parse_amount(red.amount))
        partner = None
        for pos in positives:
            if pos.id in used_positive or pos.buyer_name != red.buyer_name:
                continue
            if abs(parse_amount(pos.amount)) == red_abs:
                partner = pos
                break
        if partner is not None:
            red.paired_item_id = partner.id
            partner.paired_item_id = red.id
            used_positive.add(partner.id)
            matched += 1
    return matched, len(reds) - matched


# ---------------- 上传与批次创建 ----------------

def _validate_uploads(uploads: list[tuple[str, bytes]]) -> None:
    """上传校验：扩展名白名单 + 大小限制 + 文件数限制；违规抛 ValueError（路由层转 400）。"""
    if not uploads:
        raise ValueError("未接收到文件")
    if len(uploads) > MAX_FILES_PER_BATCH:
        raise ValueError(f"单批次文件数超过 {MAX_FILES_PER_BATCH} 上限")
    total = 0
    zip_count = 0
    for fname, content in uploads:
        safe = Path(fname or "").name
        if safe in ("", ".", ".."):
            raise ValueError("存在非法文件名")
        suffix = Path(safe).suffix.lower()
        if suffix not in ALLOWED_SUFFIXES:
            raise ValueError(f"不支持的文件类型：{safe}（仅允许 .pdf / .zip）")
        size = len(content)
        if size == 0:
            raise ValueError(f"空文件：{safe}")
        if size > MAX_FILE_BYTES:
            raise ValueError(f"文件 {safe} 超过 20MB 上限")
        total += size
        if suffix == ".zip":
            zip_count += 1
    if zip_count and len(uploads) > 1:
        raise ValueError("ZIP 上传时仅允许单个 ZIP 文件")
    if total > MAX_BATCH_BYTES:
        raise ValueError("批次总大小超过 100MB 上限")


def _unique_path(path: Path) -> Path:
    """同名冲突时追加 _1/_2 后缀（同引擎 ensure_unique_filepath）。"""
    if not path.exists():
        return path
    stem, suffix = path.stem, path.suffix
    counter = 1
    while True:
        candidate = path.with_name(f"{stem}_{counter}{suffix}")
        if not candidate.exists():
            return candidate
        counter += 1


def _save_pdf(inbox: Path, source_name: str, content: bytes) -> Path:
    """保存单个 PDF 到批次 inbox（文件名净化 + 防路径穿越 + 同名去重）。"""
    safe = Path(source_name).name
    if safe in ("", ".", ".."):
        raise ValueError("非法的文件名")
    dest = _unique_path(inbox / safe)
    if dest.resolve().parent != inbox.resolve():
        raise ValueError("非法的文件名")
    dest.write_bytes(content)
    return dest


def _extract_zip_pdfs(inbox: Path, zip_name: str, content: bytes) -> list[tuple[str, Path]]:
    """解包 ZIP 中的 PDF（拍平目录只保留文件名）；路径穿越成员直接整包拒绝。"""
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile:
        raise ValueError(f"ZIP 文件损坏：{zip_name}")
    results: list[tuple[str, Path]] = []
    with zf:
        members = [m for m in zf.infolist() if not m.is_dir() and m.filename.lower().endswith(".pdf")]
        if not members:
            raise ValueError("ZIP 内未找到 PDF 文件")
        if len(members) > MAX_FILES_PER_BATCH:
            raise ValueError(f"ZIP 内 PDF 数量超过 {MAX_FILES_PER_BATCH} 上限")
        total_unc = 0
        for member in members:
            norm = member.filename.replace("\\", "/")
            parts = [p for p in norm.split("/") if p not in ("", ".")]
            if not parts or any(p == ".." for p in parts) or norm.startswith("/") or re.match(r"^[A-Za-z]:", norm):
                raise ValueError(f"ZIP 内存在非法路径：{member.filename}")
            total_unc += member.file_size
            if total_unc > MAX_BATCH_BYTES:
                raise ValueError("ZIP 解压后总大小超过 100MB 上限")
            data = zf.read(member)
            stored = _save_pdf(inbox, parts[-1], data)
            results.append((f"{zip_name}:{member.filename}"[:255], stored))
    return results


def create_batch(db: Session, uploads: list[tuple[str, bytes]], name: str = "") -> InvoiceBatch:
    """建批次并解析：uploaded → parsed。

    引擎不可用在任何落盘/建库动作之前抛 EngineUnavailableError（503 降级，不留半成品批次）；
    上传校验失败抛 ValueError（400）；单张解析失败只标记该明细，不影响整批。
    """
    _require_engine()
    _validate_uploads(uploads)

    batch = InvoiceBatch(
        name=(name or "").strip()[:100] or f"批次{datetime.now():%Y%m%d-%H%M%S}",
        status=InvoiceBatch.STATUS_UPLOADED,
        source_files=json.dumps([Path(f).name for f, _ in uploads], ensure_ascii=False),
    )
    db.add(batch)
    db.commit()
    db.refresh(batch)

    dirs = _ensure_dirs()
    inbox = dirs["inbox"] / f"batch_{batch.id}"
    inbox.mkdir(parents=True, exist_ok=True)

    try:
        stored: list[tuple[str, Path]] = []
        for fname, content in uploads:
            if Path(fname).suffix.lower() == ".zip":
                stored.extend(_extract_zip_pdfs(inbox, Path(fname).name, content))
            else:
                stored.append((Path(fname).name, _save_pdf(inbox, fname, content)))
        batch.file_count = len(stored)

        items: list[InvoiceItem] = []
        for source_name, pdf_path in stored:
            parsed = _parse_pdf_with_engine(pdf_path)
            if parsed is None:
                item = InvoiceItem(
                    batch_id=batch.id,
                    source_file=source_name[:255],
                    stored_path=str(pdf_path)[:500],
                    status=InvoiceItem.STATUS_FAILED,
                    extraction_success=False,
                )
            else:
                item = InvoiceItem(
                    batch_id=batch.id,
                    source_file=source_name[:255],
                    stored_path=str(pdf_path)[:500],
                    invoice_number=parsed.invoice_number,
                    check_code=parsed.check_code,
                    buyer_name=parsed.buyer_name,
                    seller_name=parsed.seller_name,
                    amount=parsed.amount,
                    tax_amount=parsed.tax_amount,
                    tax_rate=parsed.tax_rate,
                    invoice_date=parsed.invoice_date,
                    invoice_type=parsed.invoice_type,
                    invoice_direction=parsed.invoice_direction,
                    project_name=parsed.project_name,
                    brand_name=parsed.brand_name,
                    remark=parsed.remark,
                    extraction_success=parsed.extraction_success,
                    status=InvoiceItem.STATUS_PARSED if parsed.extraction_success else InvoiceItem.STATUS_INCOMPLETE,
                )
            db.add(item)
            items.append(item)
        db.commit()
        for item in items:
            db.refresh(item)

        duplicate_count = mark_duplicates(items, existing_dedup_keys(db, exclude_batch_id=batch.id))
        red_matched, red_unmatched = pair_red_invoices(items)

        batch.item_count = len(items)
        batch.parsed_count = sum(1 for i in items if i.status in _ACTIVE_ITEM_STATUS)
        batch.failed_count = sum(1 for i in items if i.status == InvoiceItem.STATUS_FAILED)
        batch.duplicate_count = duplicate_count
        batch.red_matched_count = red_matched
        batch.red_unmatched_count = red_unmatched
        batch.total_amount = round(sum(parse_amount(i.amount) for i in items if i.status in _ACTIVE_ITEM_STATUS), 2)
        batch.status = InvoiceBatch.STATUS_PARSED
        db.commit()
        db.refresh(batch)
        return batch
    except Exception as exc:
        logger.exception("批次 %s 解析流程异常", batch.id)
        batch.error = f"{type(exc).__name__}: {exc}"[:500]
        db.commit()
        raise


# ---------------- 查询与人工修正 ----------------

def list_batches(db: Session, limit: int = 50) -> list[InvoiceBatch]:
    """批次列表（按 id 倒序）。"""
    limit = max(1, min(int(limit), 200))
    return db.query(InvoiceBatch).order_by(InvoiceBatch.id.desc()).limit(limit).all()


def get_batch(db: Session, batch_id: int) -> Optional[InvoiceBatch]:
    return db.query(InvoiceBatch).filter(InvoiceBatch.id == batch_id).first()


def get_item(db: Session, item_id: int) -> Optional[InvoiceItem]:
    return db.query(InvoiceItem).filter(InvoiceItem.id == item_id).first()


def patch_item(
    db: Session,
    item: InvoiceItem,
    project_name: Optional[str] = None,
    remark: Optional[str] = None,
) -> InvoiceItem:
    """人工修正项目归类 / 备注；有实际改动时打 corrected 标记。"""
    changed = False
    if project_name is not None:
        item.project_name = project_name.strip()[:100]
        changed = True
    if remark is not None:
        item.remark = remark.strip()[:500]
        changed = True
    if changed:
        item.corrected = True
        db.commit()
        db.refresh(item)
    return item


# ---------------- 导出与下载 ----------------

def _write_project_list(project_dir: Path, project: str, entries: list[dict]) -> Path:
    """生成单项目「发票清单.xlsx」（格式镜像引擎 ExcelGenerator.generate_invoice_list）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = project[:31] or "项目清单"
    headers = ["序号", "票号", "购方名称", "发票金额", "发票类型", "备注"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    total_amount = 0.0
    special_count = 0
    general_count = 0
    for row, entry in enumerate(entries, 2):
        item: InvoiceItem = entry["item"]
        ws.cell(row=row, column=1, value=row - 1)
        ws.cell(row=row, column=2, value=item.invoice_number)
        ws.cell(row=row, column=3, value=item.buyer_name)
        amount = parse_amount(item.amount)
        amount_cell = ws.cell(row=row, column=4, value=amount)
        amount_cell.number_format = _AMOUNT_NUMBER_FORMAT
        total_amount += amount
        ws.cell(row=row, column=5, value=item.invoice_type)
        if item.invoice_type == "专票":
            special_count += 1
        elif item.invoice_type == "普票":
            general_count += 1
        ws.cell(row=row, column=6, value=item.remark)
    stat_row = len(entries) + 3
    ws.cell(row=stat_row, column=1, value="统计信息")
    stats = [("发票总额", round(total_amount, 2)), ("发票张数", len(entries)), ("专票张数", special_count), ("普票张数", general_count)]
    for offset, (label, value) in enumerate(stats):
        ws.cell(row=stat_row + offset, column=2, value=label)
        ws.cell(row=stat_row + offset, column=3, value=value)
        for col in range(1, 4):
            ws.cell(row=stat_row + offset, column=col).font = Font(bold=True)
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[column[0].column_letter].width = min(width + 2, 50)
    path = project_dir / "发票清单.xlsx"
    wb.save(path)
    return path


def _write_summary_list(batch_dir: Path, groups: dict[str, list[dict]]) -> Path:
    """生成批次「汇总清单.xlsx」（格式镜像引擎 generate_summary_list）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总清单"
    headers = ["序号", "项目名称", "发票张数", "发票总额", "专票张数", "普票张数"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    total_rows = 0
    total_amount = 0.0
    total_special = 0
    total_general = 0
    for row, (project, entries) in enumerate(groups.items(), 2):
        amounts = [parse_amount(e["item"].amount) for e in entries]
        special = sum(1 for e in entries if e["item"].invoice_type == "专票")
        general = sum(1 for e in entries if e["item"].invoice_type == "普票")
        ws.cell(row=row, column=1, value=row - 1)
        ws.cell(row=row, column=2, value=project)
        ws.cell(row=row, column=3, value=len(entries))
        amount_cell = ws.cell(row=row, column=4, value=round(sum(amounts), 2))
        amount_cell.number_format = _AMOUNT_NUMBER_FORMAT
        ws.cell(row=row, column=5, value=special)
        ws.cell(row=row, column=6, value=general)
        total_rows += len(entries)
        total_amount += sum(amounts)
        total_special += special
        total_general += general
    total_row = len(groups) + 2
    ws.cell(row=total_row, column=2, value="合计").font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=total_rows).font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=4, value=round(total_amount, 2))
    total_cell.number_format = _AMOUNT_NUMBER_FORMAT
    total_cell.font = Font(bold=True)
    ws.cell(row=total_row, column=5, value=total_special).font = Font(bold=True)
    ws.cell(row=total_row, column=6, value=total_general).font = Font(bold=True)
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[column[0].column_letter].width = min(width + 2, 50)
    path = batch_dir / "汇总清单.xlsx"
    wb.save(path)
    return path


def export_batch(db: Session, batch: InvoiceBatch) -> dict:
    """导出批次：按重命名规则复制 PDF → 项目目录，生成清单 xlsx，打包 ZIP。

    可重复导出（先清空本批次旧导出目录）；无可导出明细抛 ValueError。
    """
    items = (
        db.query(InvoiceItem)
        .filter(InvoiceItem.batch_id == batch.id, InvoiceItem.status.in_(_ACTIVE_ITEM_STATUS))
        .order_by(InvoiceItem.id.asc())
        .all()
    )
    if not items:
        raise ValueError("批次内无可导出的发票明细")

    dirs = _ensure_dirs()
    outroot = dirs["outbox"]
    batch_dir = outroot / f"batch_{batch.id}"
    if batch_dir.exists():
        shutil.rmtree(batch_dir)  # 仅清理本模块生成的本批次导出目录
    batch_dir.mkdir(parents=True, exist_ok=True)

    # 按项目分组成 {项目: [{item, dest}]}
    groups: dict[str, list[dict]] = {}
    for item in items:
        project = _NON_WORD_CHARS.sub("", sanitize_filename(item.project_name.strip() or "未知项目"))[:30] or "未知项目"
        groups.setdefault(project, []).append({"item": item, "dest": None})

    for project, entries in groups.items():
        for entry in entries:
            item: InvoiceItem = entry["item"]
            ym = parse_date_to_year_month(item.invoice_date)
            dir_name = f"{ym}{project}" if ym else project
            project_dir = batch_dir / dir_name
            project_dir.mkdir(parents=True, exist_ok=True)
            src = Path(item.stored_path)
            if src.is_file():
                dest = _unique_path(project_dir / build_renamed_filename(item))
                shutil.copy2(src, dest)
                entry["dest"] = dest
            else:
                logger.warning("批次 %s 明细 %s 的入库文件缺失：%s", batch.id, item.id, src)

    for project, entries in groups.items():
        first_dest = next((e["dest"] for e in entries if e["dest"] is not None), None)
        if first_dest is not None:
            _write_project_list(first_dest.parent, project, entries)
    _write_summary_list(batch_dir, groups)

    zip_name = f"batch_{batch.id}_export.zip"
    zip_path = outroot / zip_name
    if zip_path.exists():
        zip_path.unlink()  # 重新导出覆盖旧包
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for file in sorted(batch_dir.rglob("*")):
            if file.is_file():
                zf.write(file, file.relative_to(batch_dir).as_posix())

    batch.export_file = zip_path.name
    batch.status = InvoiceBatch.STATUS_EXPORTED
    db.commit()
    db.refresh(batch)
    return {
        "ok": True,
        "batch_id": batch.id,
        "status": batch.status,
        "export_file": batch.export_file,
        "zip_size": zip_path.stat().st_size,
        "projects": len(groups),
    }


def export_zip_path(batch: InvoiceBatch) -> Optional[Path]:
    """导出 ZIP 的服务端路径（限制在模块 outbox 内，防路径穿越）；不存在返回 None。"""
    if not batch.export_file:
        return None
    outroot = _ensure_dirs()["outbox"].resolve()
    path = (outroot / Path(batch.export_file).name).resolve()
    if path.parent != outroot or not path.is_file():
        return None
    return path


# ---------------- 响应组装 ----------------

def batch_to_dict(batch: InvoiceBatch) -> dict:
    """批次 ORM → API 输出 dict（source_files JSON 文本转数组）。"""
    try:
        source_files = json.loads(batch.source_files or "[]")
    except (ValueError, TypeError):
        source_files = []
    return {
        "id": batch.id,
        "name": batch.name,
        "status": batch.status,
        "file_count": batch.file_count,
        "item_count": batch.item_count,
        "parsed_count": batch.parsed_count,
        "failed_count": batch.failed_count,
        "duplicate_count": batch.duplicate_count,
        "red_matched_count": batch.red_matched_count,
        "red_unmatched_count": batch.red_unmatched_count,
        "total_amount": batch.total_amount,
        "source_files": source_files,
        "export_file": batch.export_file,
        "error": batch.error,
        "created_at": batch.created_at,
        "updated_at": batch.updated_at,
    }


def item_to_dict(item: InvoiceItem, partner: Optional[InvoiceItem] = None) -> dict:
    """明细 ORM → API 输出 dict（stored_path 不对外）。"""
    return {
        "id": item.id,
        "batch_id": item.batch_id,
        "source_file": item.source_file,
        "invoice_number": item.invoice_number,
        "check_code": item.check_code,
        "buyer_name": item.buyer_name,
        "seller_name": item.seller_name,
        "amount": item.amount,
        "tax_amount": item.tax_amount,
        "tax_rate": item.tax_rate,
        "invoice_date": item.invoice_date,
        "invoice_type": item.invoice_type,
        "invoice_direction": item.invoice_direction,
        "project_name": item.project_name,
        "brand_name": item.brand_name,
        "remark": item.remark,
        "extraction_success": item.extraction_success,
        "corrected": item.corrected,
        "status": item.status,
        "paired_item_id": item.paired_item_id,
        "paired_invoice_number": partner.invoice_number if partner else "",
        "created_at": item.created_at,
        "updated_at": item.updated_at,
    }


def batch_detail_dict(db: Session, batch: InvoiceBatch) -> dict:
    """批次详情（含明细；配对对方号码一次性查出，避免 N+1）。"""
    items = db.query(InvoiceItem).filter(InvoiceItem.batch_id == batch.id).order_by(InvoiceItem.id.asc()).all()
    partner_ids = {i.paired_item_id for i in items if i.paired_item_id}
    partners = {}
    if partner_ids:
        partners = {
            p.id: p
            for p in db.query(InvoiceItem).filter(InvoiceItem.id.in_(partner_ids)).all()
        }
    data = batch_to_dict(batch)
    data["items"] = [item_to_dict(i, partners.get(i.paired_item_id)) for i in items]
    return data


# ---------------- Agent 工具 ----------------

def invoice_batch_stats() -> dict:
    """Agent 工具（readonly）：发票处理模块统计概览。"""
    try:
        with SessionLocal() as db:
            total_batches = db.query(func.count(InvoiceBatch.id)).scalar() or 0
            by_status = dict(
                db.query(InvoiceBatch.status, func.count()).group_by(InvoiceBatch.status).all()
            )
            total_items = db.query(func.count(InvoiceItem.id)).scalar() or 0
            duplicates = (
                db.query(func.count(InvoiceItem.id)).filter(InvoiceItem.status == InvoiceItem.STATUS_DUPLICATE).scalar() or 0
            )
            failed = (
                db.query(func.count(InvoiceItem.id)).filter(InvoiceItem.status == InvoiceItem.STATUS_FAILED).scalar() or 0
            )
            red_matched = (
                db.query(func.count(InvoiceItem.id))
                .filter(InvoiceItem.invoice_direction == "冲红发票", InvoiceItem.paired_item_id.isnot(None))
                .scalar()
                or 0
            )
            red_unmatched = (
                db.query(func.count(InvoiceItem.id))
                .filter(InvoiceItem.invoice_direction == "冲红发票", InvoiceItem.paired_item_id.is_(None))
                .scalar()
                or 0
            )
            net = db.query(func.coalesce(func.sum(InvoiceBatch.total_amount), 0.0)).scalar() or 0.0
        return {
            "ok": True,
            "module": "finance_invoice",
            "batches": total_batches,
            "batch_by_status": by_status,
            "items": total_items,
            "duplicates": duplicates,
            "failed": failed,
            "red_matched": red_matched,
            "red_unmatched": red_unmatched,
            "net_amount": round(float(net), 2),
        }
    except Exception as exc:
        logger.exception("invoice_batch_stats 执行失败")
        return {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
