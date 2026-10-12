"""finance_invoice 模块测试：合成数据 + mock 解析器，不触碰真实业务文件、不写 tmp_path 之外的位置。

覆盖：
- 文本字段解析（真实引擎 + 示例配置直喂合成文本，引擎缺失自动 skip）与销方名补充提取；
- 去重幂等键 / 批内 + 跨批次去重 / 冲红一对一配对（纯逻辑，ORM  detached 对象）；
- 上传建批次状态机 uploaded→parsed→exported（monkeypatch 引擎解析器）、ZIP 解包拍平与路径穿越拦截、
  扩展名/大小/ZIP 混传校验；
- 引擎缺失时 POST /batches 503 降级（不留半成品批次）；
- PATCH 人工修正、导出清单 xlsx + ZIP、下载与未导出 409；
- invoice_batch_stats 只读 Agent 工具。
"""
import asyncio
import io
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.config import settings
from app.database import engine
from app.modules.finance_invoice import engine as engine_mod
from app.modules.finance_invoice import router as invoice_router
from app.modules.finance_invoice import service
from app.modules.finance_invoice.models import InvoiceBatch, InvoiceItem

# 合成的假 PDF 内容：解析器全部被 mock，字节不会被真正解析
FAKE_PDF = b"%PDF-1.4 synthetic-test-bytes"


def make_fields(**kw) -> dict:
    """构造一份 ParsedInvoice 字段（全为虚构示例值）。"""
    base = dict(
        invoice_number="10000001",
        check_code="c001",
        buyer_name="示例购方公司",
        seller_name="示例销方公司",
        amount="100.00",
        tax_amount="13.00",
        tax_rate="0.13",
        invoice_date="2026年10月01日",
        invoice_type="普票",
        invoice_direction="正数发票",
        project_name="示例项目",
        brand_name="",
        remark="合成备注",
        extraction_success=True,
    )
    base.update(kw)
    return base


@pytest.fixture()
def finance_dir(tmp_path, monkeypatch):
    """模块数据目录重定向到 tmp_path（不碰真实 finance 数据目录）。"""
    monkeypatch.setattr(settings, "finance_data_dir", str(tmp_path / "finance"))
    return tmp_path / "finance"


@pytest.fixture()
def mock_engine(monkeypatch):
    """打桩引擎：_require_engine 放行；按文件名返回罐头解析结果（未登记的文件返回 None=解析失败）。"""
    canned: dict[str, dict] = {}
    monkeypatch.setattr(service, "_require_engine", lambda: None)

    def fake_parse(pdf_path):
        fields = canned.get(Path(pdf_path).name)
        return service.ParsedInvoice(**fields) if fields is not None else None

    monkeypatch.setattr(service, "_parse_pdf_with_engine", fake_parse)
    return canned


def upload(client, files: list[tuple[str, bytes]], name: str = ""):
    """files: [(filename, content), ...]；走 POST /batches multipart。"""
    multipart = [("files", (fname, content, "application/pdf")) for fname, content in files]
    data = {"name": name} if name else {}
    return client.post("/api/finance/invoice/batches", files=multipart, data=data)


# ---------------- 模块注册与引擎自检 ----------------

def test_module_registered_and_tables_created(client):
    from app.modules import get_loaded_modules

    loaded = {m.name: m for m in get_loaded_modules()}
    assert "finance_invoice" in loaded
    manifest = loaded["finance_invoice"]
    assert manifest.prefix == "/api/finance/invoice"
    assert manifest.tables == ["mod_invoice_batch", "mod_invoice_item"]
    assert manifest.menu == {"group": "财务", "icon": "Tickets", "order": 20}
    with engine.connect() as conn:
        names = {r[0] for r in conn.execute(text("SELECT name FROM sqlite_master WHERE type='table'"))}
    assert {"mod_invoice_batch", "mod_invoice_item"} <= names


def test_engine_status_endpoint_shape(client):
    resp = client.get("/api/finance/invoice/engine/status")
    assert resp.status_code == 200
    body = resp.json()
    assert body["module"] == "finance_invoice"
    assert isinstance(body["available"], bool)
    assert isinstance(body["root"], str) and body["root"]


def test_create_batch_handler_remains_coroutine():
    """create_batch 保持协程（逐文件 await f.read 语义不变），PDF 解析重活经 asyncio.to_thread 进线程池。"""
    assert asyncio.iscoroutinefunction(invoice_router.create_batch)


# ---------------- 文本解析（真实引擎，缺失则 skip） ----------------

def _real_processor(monkeypatch):
    """加载真实引擎并强制示例配置（INVOICE_CONFIG_EXAMPLE=1 跳过引擎目录的本机 local 配置）。"""
    root = engine_mod.default_engine_root()
    if not (root / "invoice_processor" / "src" / "pdf_processor.py").is_file():
        pytest.skip("发票引擎不在默认位置")
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    monkeypatch.setenv("INVOICE_CONFIG_EXAMPLE", "1")
    try:
        from invoice_processor.src.config_manager import ConfigManager
        from invoice_processor.src.pdf_processor import PDFInvoiceProcessor
    except Exception as exc:
        pytest.skip(f"发票引擎不可导入：{exc}")
    return PDFInvoiceProcessor(ConfigManager())


def test_parse_text_with_real_engine(monkeypatch):
    processor = _real_processor(monkeypatch)
    text = (
        "电子发票（增值税专用发票） 发票号码：26937000000123 校验码：12345678901234567890 "
        "开票日期：2026年10月01日 购 名称：示例科技有限公司 销 名称：云岭商贸有限公司 "
        "价税合计 ￥113.00 税率：13% 备注：办公用品采购"
    )
    parsed = service.parse_invoice_text(processor, text)
    assert parsed.invoice_number == "26937000000123"
    assert parsed.buyer_name == "示例科技有限公司"
    assert parsed.seller_name == "云岭商贸有限公司"  # 引擎不解析销方，本模块补齐
    assert parsed.amount == "113.00"
    assert parsed.tax_amount == "13.00"  # 引擎按税率价外折算
    assert parsed.invoice_date == "2026年10月01日"
    assert parsed.invoice_type == "专票"
    assert parsed.invoice_direction == "正数发票"
    assert parsed.project_name == "办公用品采购"  # 引擎回退：取备注前 10 字
    assert parsed.extraction_success is True


def test_parse_text_red_invoice_with_real_engine(monkeypatch):
    processor = _real_processor(monkeypatch)
    text = (
        "电子发票（普通发票） 发票号码：26937000000999 开票日期：2026年09月15日 "
        "购 名称：示例科技有限公司 销 名称：云岭商贸有限公司 价税合计 ￥-113.00 冲红"
    )
    parsed = service.parse_invoice_text(processor, text)
    assert parsed.invoice_direction == "冲红发票"
    assert service.parse_amount(parsed.amount) == -113.00
    assert parsed.invoice_type == "普票"


def test_extract_seller_name_patterns():
    assert service.extract_seller_name("购 名称：甲公司 销 名称：乙有限责任公司 价税合计") == "乙有限责任公司"
    assert service.extract_seller_name("销售方 名称：丙供应商") == "丙供应商"
    assert service.extract_seller_name("没有任何销方字段") == ""
    # 命中停用词时不应采信
    assert service.extract_seller_name("销 名称：发票号码") == ""


# ---------------- 去重与冲红配对（纯逻辑） ----------------

def test_dedup_key_rule():
    assert service.dedup_key("123", "ABC") == "123|ABC"
    assert service.dedup_key("123", "") == "123|"
    assert service.dedup_key("", "x") == ""
    assert service.dedup_key("未知", "") == ""


def test_mark_duplicates_within_batch():
    a = InvoiceItem(invoice_number="1001", check_code="c1", status=InvoiceItem.STATUS_PARSED)
    b = InvoiceItem(invoice_number="1001", check_code="c1", status=InvoiceItem.STATUS_PARSED)
    c = InvoiceItem(invoice_number="1002", check_code="", status=InvoiceItem.STATUS_PARSED)
    f = InvoiceItem(invoice_number="1001", check_code="c1", status=InvoiceItem.STATUS_FAILED)
    dup = service.mark_duplicates([a, b, c, f], existing_keys=set())
    assert dup == 1
    assert a.status == InvoiceItem.STATUS_PARSED
    assert b.status == InvoiceItem.STATUS_DUPLICATE  # 后来者被标记
    assert c.status == InvoiceItem.STATUS_PARSED
    assert f.status == InvoiceItem.STATUS_FAILED  # 失败明细不参与去重


def test_mark_duplicates_against_existing_keys():
    item = InvoiceItem(invoice_number="1001", check_code="c1", status=InvoiceItem.STATUS_PARSED)
    assert service.mark_duplicates([item], existing_keys={"1001|c1"}) == 1
    assert item.status == InvoiceItem.STATUS_DUPLICATE


def test_pair_red_invoices_one_to_one():
    pos = InvoiceItem(id=1, buyer_name="示例购方", amount="113.00", invoice_direction="正数发票", status="parsed")
    red = InvoiceItem(id=2, buyer_name="示例购方", amount="-113.00", invoice_direction="冲红发票", status="parsed")
    red2 = InvoiceItem(id=3, buyer_name="示例购方", amount="-113.00", invoice_direction="冲红发票", status="parsed")
    matched, unmatched = service.pair_red_invoices([pos, red, red2])
    assert (matched, unmatched) == (1, 1)  # 一对一：一张原票只能配一张红票
    assert red.paired_item_id == 1 and pos.paired_item_id == 2
    assert red2.paired_item_id is None


def test_pair_red_invoices_buyer_mismatch_and_duplicate_excluded():
    pos = InvoiceItem(id=1, buyer_name="购方甲", amount="100.00", invoice_direction="正数发票", status="parsed")
    red_other = InvoiceItem(id=2, buyer_name="购方乙", amount="-100.00", invoice_direction="冲红发票", status="parsed")
    red_dup = InvoiceItem(id=3, buyer_name="购方甲", amount="-100.00", invoice_direction="冲红发票", status="duplicate")
    matched, unmatched = service.pair_red_invoices([pos, red_other, red_dup])
    assert (matched, unmatched) == (0, 1)  # 购方不同不配；重复红票不参与也不计入未配对
    assert pos.paired_item_id is None and red_other.paired_item_id is None


# ---------------- 上传与状态机（mock 解析器） ----------------

def test_upload_pdfs_state_machine(client, finance_dir, mock_engine):
    mock_engine["a.pdf"] = make_fields(invoice_number="8001", check_code="c1")
    mock_engine["b.pdf"] = make_fields(invoice_number="8001", check_code="c1")  # 与 a 同号同校验码 → 重复
    mock_engine["red.pdf"] = make_fields(invoice_number="8002", amount="-100.00", invoice_direction="冲红发票")
    # c.pdf 未登记 → 解析失败
    resp = upload(client, [("a.pdf", FAKE_PDF), ("b.pdf", FAKE_PDF), ("red.pdf", FAKE_PDF), ("c.pdf", FAKE_PDF)], name="十月票")
    assert resp.status_code == 201
    body = resp.json()
    assert body["name"] == "十月票"
    assert body["status"] == "parsed"  # uploaded → parsed
    assert (body["file_count"], body["item_count"]) == (4, 4)
    assert (body["parsed_count"], body["duplicate_count"], body["failed_count"]) == (2, 1, 1)
    assert (body["red_matched_count"], body["red_unmatched_count"]) == (1, 0)
    assert body["total_amount"] == 0.0  # 100 + (-100)
    items = {i["source_file"]: i for i in body["items"]}
    assert items["b.pdf"]["status"] == "duplicate"
    assert items["c.pdf"]["status"] == "failed"
    assert items["red.pdf"]["paired_item_id"] == items["a.pdf"]["id"]
    assert items["red.pdf"]["paired_invoice_number"] == "8001"
    assert items["a.pdf"]["seller_name"] == "示例销方公司"
    assert "stored_path" not in items["a.pdf"]  # 服务端内部路径不对外
    # PDF 落盘在 tmp 指向的模块 inbox，不碰真实目录
    stored = finance_dir / "finance_invoice" / "inbox" / f"batch_{body['id']}" / "a.pdf"
    assert stored.is_file() and stored.read_bytes() == FAKE_PDF


def test_upload_cross_batch_duplicate(client, finance_dir, mock_engine):
    mock_engine["a.pdf"] = make_fields(invoice_number="8001", check_code="c1")
    assert upload(client, [("a.pdf", FAKE_PDF)]).status_code == 201
    # 第二批再次上传同号同校验码 → 跨批次去重
    resp = upload(client, [("a.pdf", FAKE_PDF)])
    assert resp.status_code == 201
    body = resp.json()
    assert body["duplicate_count"] == 1 and body["parsed_count"] == 0
    assert body["items"][0]["status"] == "duplicate"


def test_upload_zip_extract_and_flatten(client, finance_dir, mock_engine):
    mock_engine["x.pdf"] = make_fields(invoice_number="9001")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("子目录/x.pdf", FAKE_PDF)
        zf.writestr("y.pdf", FAKE_PDF)  # 未登记 → failed
        zf.writestr("说明.txt", b"not a pdf")  # 非 PDF 成员跳过
    resp = upload(client, [("pack.zip", buf.getvalue())])
    assert resp.status_code == 201
    body = resp.json()
    assert body["file_count"] == 2 and body["item_count"] == 2
    by_src = {i["source_file"]: i for i in body["items"]}
    assert by_src["pack.zip:子目录/x.pdf"]["status"] == "parsed"
    assert by_src["pack.zip:y.pdf"]["status"] == "failed"
    inbox = finance_dir / "finance_invoice" / "inbox" / f"batch_{body['id']}"
    assert (inbox / "x.pdf").is_file()  # 目录被拍平


def test_upload_zip_path_traversal_rejected(client, finance_dir, mock_engine):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../evil.pdf", FAKE_PDF)
    resp = upload(client, [("evil.zip", buf.getvalue())])
    assert resp.status_code == 400
    assert "非法路径" in resp.json()["detail"]


def test_upload_validation(client, finance_dir, mock_engine):
    # 扩展名白名单
    resp = upload(client, [("a.txt", b"hi")])
    assert resp.status_code == 400 and "不支持的文件类型" in resp.json()["detail"]
    # ZIP 必须单传
    resp = upload(client, [("a.pdf", FAKE_PDF), ("b.zip", b"PK\x03\x04")])
    assert resp.status_code == 400 and "ZIP" in resp.json()["detail"]


def test_upload_oversize_rejected(client, finance_dir, mock_engine, monkeypatch):
    monkeypatch.setattr(service, "MAX_FILE_BYTES", 8)
    resp = upload(client, [("big.pdf", b"x" * 9)])
    assert resp.status_code == 400 and "上限" in resp.json()["detail"]


def test_upload_503_when_engine_missing(client, finance_dir, monkeypatch, tmp_path):
    """引擎路径无效 → 503 降级，且不留半成品批次。"""
    monkeypatch.setattr(settings, "finance_invoice_engine_path", str(tmp_path / "no_such_engine"))
    engine_mod.reset_engine_cache()
    try:
        resp = upload(client, [("a.pdf", FAKE_PDF)])
        assert resp.status_code == 503
        assert "引擎不可用" in resp.json()["detail"]
        assert client.get("/api/finance/invoice/batches").json() == []
        status = client.get("/api/finance/invoice/engine/status").json()
        assert status["available"] is False
    finally:
        engine_mod.reset_engine_cache()


def test_batch_detail_404(client):
    assert client.get("/api/finance/invoice/batches/9999").status_code == 404


# ---------------- 人工修正 ----------------

def test_patch_item(client, finance_dir, mock_engine):
    mock_engine["a.pdf"] = make_fields()
    batch = upload(client, [("a.pdf", FAKE_PDF)]).json()
    item_id = batch["items"][0]["id"]
    resp = client.patch(f"/api/finance/invoice/items/{item_id}", json={"project_name": "修正后项目", "remark": "人工备注"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["project_name"] == "修正后项目" and body["remark"] == "人工备注"
    assert body["corrected"] is True
    assert client.patch("/api/finance/invoice/items/9999", json={"remark": "x"}).status_code == 404


# ---------------- 导出与下载 ----------------

def _exported_batch(client, mock_engine) -> dict:
    mock_engine["a.pdf"] = make_fields(invoice_number="8001", amount="100.00", invoice_type="专票")
    mock_engine["b.pdf"] = make_fields(invoice_number="8002", amount="50.00", project_name="另一项目")
    batch = upload(client, [("a.pdf", FAKE_PDF), ("b.pdf", FAKE_PDF)]).json()
    resp = client.post(f"/api/finance/invoice/batches/{batch['id']}/export")
    assert resp.status_code == 200
    result = resp.json()
    assert result["ok"] and result["projects"] == 2 and result["zip_size"] > 0
    return client.get(f"/api/finance/invoice/batches/{batch['id']}").json()


def test_export_generates_xlsx_and_zip(client, finance_dir, mock_engine, tmp_path):
    batch = _exported_batch(client, mock_engine)
    assert batch["status"] == "exported"  # parsed → exported
    zip_path = finance_dir / "finance_invoice" / "outbox" / batch["export_file"]
    assert zip_path.is_file()
    with zipfile.ZipFile(zip_path) as zf:
        names = zf.namelist()
        # 引擎重命名规则：{购方}{金额}{类型}.pdf，目录 {年月}{项目}
        assert "26年10月示例项目/示例购方公司100.00专票.pdf" in names
        assert any(name.endswith("发票清单.xlsx") for name in names)
        assert "汇总清单.xlsx" in names
        zf.extract("汇总清单.xlsx", tmp_path)
    wb = load_workbook(tmp_path / "汇总清单.xlsx")
    ws = wb.active
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=4)]
    assert rows[0] == ["序号", "项目名称", "发票张数", "发票总额", "专票张数", "普票张数"]
    assert rows[3][1] == "合计" and rows[3][2] == 2 and rows[3][3] == 150.0


def test_export_empty_batch_400(client, finance_dir, mock_engine):
    # c.pdf 未登记罐头 → 全部解析失败 → 无可导出明细
    batch = upload(client, [("c.pdf", FAKE_PDF)]).json()
    resp = client.post(f"/api/finance/invoice/batches/{batch['id']}/export")
    assert resp.status_code == 400 and "无可导出" in resp.json()["detail"]


def test_download_zip(client, finance_dir, mock_engine):
    batch = _exported_batch(client, mock_engine)
    resp = client.get(f"/api/finance/invoice/batches/{batch['id']}/download")
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"  # ZIP 魔数


def test_download_before_export_409(client, finance_dir, mock_engine):
    mock_engine["a.pdf"] = make_fields()
    batch = upload(client, [("a.pdf", FAKE_PDF)]).json()
    resp = client.get(f"/api/finance/invoice/batches/{batch['id']}/download")
    assert resp.status_code == 409


# ---------------- Agent 工具 ----------------

def test_invoice_batch_stats_tool(client, finance_dir, mock_engine):
    from app.agents.tools import TOOL_SPECS

    assert "invoice_batch_stats" in TOOL_SPECS
    spec = TOOL_SPECS["invoice_batch_stats"]
    assert spec["readonly"] is True
    mock_engine["a.pdf"] = make_fields()
    upload(client, [("a.pdf", FAKE_PDF)])
    result = spec["func"]()
    assert result["ok"] is True
    assert result["module"] == "finance_invoice"
    assert result["batches"] >= 1 and result["items"] >= 1
    assert result["net_amount"] == 100.0
