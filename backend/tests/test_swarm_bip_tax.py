"""Swarm 任务 bip-tax-import 聚焦测试：BIP 预览/入库与税务两条财务报表导入路径。

覆盖（对应审计台账 A01 无界上传 / A05 同步重活阻塞 / A06 行数上限 / A08 LIKE 通配符 /
A09 异常回带本机路径）：
- 上传经共享 read_upload_limited 限额读取（超限 400、共享模块缺失明确 503 不退回无界读取）；
- BIP 预览与税务导入的慢解析经 asyncio.to_thread 卸载，慢解析期间事件循环保持响应；
- BIP/税务 Excel 在写库/产出前以可检测溢出方式限制 MAX_IMPORT_ROWS，超限无半成品；
- subject_prefix 的 % / _ / 反斜杠按字面匹配；
- 税务导入对外错误不回带本机异常路径，服务器保留完整日志；
- 二轮复核新增：BIP 预览/确认解析异常与三个模板下载端点的对外错误同样不回带本机
  异常路径（服务器保留完整日志）；行数超限/解析失败不在 tmp 遗留被拒上传文件（无半成品）；
- 正常预览→确认导入与正常模板下载/导入行为不变。

约定：
- 共享接口（app.upload.read_upload_limited / app.import_limits.MAX_IMPORT_ROWS）由 core
  任务提供；本文件用与契约一致的测试替身和 monkeypatch 小值验证消费侧行为，不生成巨型文件；
- 全部使用合成数据与临时目录（conftest 已钉住临时 SQLite/调试目录），不碰真实数据。
"""
import asyncio
import io
import logging
import openpyxl
import time

import pandas as pd
import pytest
from fastapi import UploadFile

from app.modules.finance_ledger_bip import router as bip_router
from app.modules.finance_ledger_bip import service as bip_svc
from app.modules.finance_tax.routers import financial_import as tax_router
from app.modules.finance_tax.services import financial_data_importer as tax_importer
from app.upload import UploadTooLargeError

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PARSE_SLEEP = 0.5   # 合成慢解析时长
PROBE_SLEEP = 0.1   # 事件循环探针睡眠（解析开始前排入循环）
PROBE_LIMIT = 0.25  # 探针耗时上限：低于此值说明循环未被同步解析阻塞


# ---------------- 合成数据与测试替身 ----------------

def _aux_rows() -> list[dict]:
    """辅助明细账合成数据：科目编码覆盖 1002/1122/6602 前缀（与既有 BIP 测试同口径）。"""
    return [
        {"科目编码": "100201", "科目名称": "银行存款-基本户", "日期": "2026-09-05", "核算账簿": "测试账簿甲",
         "凭证号": "记-0011", "摘要": "收到货款", "借方本币": 800.0, "贷方本币": 0.0},
        {"科目编码": "112201", "科目名称": "应收账款-客户A", "日期": "2026-09-05", "核算账簿": "测试账簿甲",
         "凭证号": "记-0011", "摘要": "收到货款", "借方本币": 0.0, "贷方本币": 800.0},
        {"科目编码": "660201", "科目名称": "管理费用-办公费", "日期": "2026-09-08", "核算账簿": "测试账簿甲",
         "凭证号": "记-0012", "摘要": "购买耗材", "借方本币": 300.0, "贷方本币": 0.0},
        {"科目编码": "100202", "科目名称": "银行存款-一般户", "日期": "2026-09-08", "核算账簿": "测试账簿甲",
         "凭证号": "记-0012", "摘要": "购买耗材", "借方本币": 0.0, "贷方本币": 300.0},
    ]


def _xlsx_bytes(rows, sheet: str = "Sheet") -> bytes:
    """pandas 合成 xlsx 字节（sheet 可指定，用于走模板/XML 两条解析路径）。"""
    buf = io.BytesIO()
    pd.DataFrame(rows).to_excel(buf, index=False, sheet_name=sheet, engine="openpyxl")
    return buf.getvalue()


def _income_rows(n: int = 4) -> list[list]:
    """利润表两列合成行（项目名称 + 金额），n 为数据行数。"""
    return [["项目名称", "本期金额（万元）"]] + [[f"项目{i}", float(i)] for i in range(n)]


async def _fake_read_upload_limited(file, max_bytes, chunk_size=1024 * 1024):
    """与 core 契约一致的测试替身：分块累积读取，超限抛真实 UploadTooLargeError。"""
    chunks, total = [], 0
    while True:
        chunk = await file.read(chunk_size)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise UploadTooLargeError(max_bytes)
        chunks.append(chunk)
    return b"".join(chunks)


def _patch_reader(monkeypatch, router_module):
    """把路由模块的共享读取器替换为契约替身（共享模块落地前后行为一致）。"""
    monkeypatch.setattr(router_module, "read_upload_limited", _fake_read_upload_limited, raising=False)


@pytest.fixture()
def bip_dir(monkeypatch, tmp_path):
    """模块数据目录重定向到 tmp_path（测试不写真实数据目录）。"""
    d = tmp_path / "bip_data"
    monkeypatch.setattr(bip_svc, "module_data_dir", lambda: d)
    return d


def _bip_preview(client, payload: bytes, name: str = "s.xlsx"):
    return client.post("/api/finance/ledger/bip/imports/preview",
                       files={"file": (name, payload, XLSX_MIME)})


# ---------------- 上传：共享限额读取 ----------------

def test_bip_preview_reads_upload_via_shared_limited_reader(client, bip_dir, monkeypatch):
    """BIP 预览经共享 read_upload_limited 读取，并传入模块既有 10MB 口径。"""
    seen: list[int] = []

    async def spy(file, max_bytes, chunk_size=1024 * 1024):
        seen.append(max_bytes)
        return await _fake_read_upload_limited(file, max_bytes, chunk_size)

    monkeypatch.setattr(bip_router, "read_upload_limited", spy, raising=False)
    resp = _bip_preview(client, _xlsx_bytes(_aux_rows()))
    assert resp.status_code == 200, resp.text
    assert seen == [bip_svc.MAX_UPLOAD_BYTES]
    assert resp.json()["file_type"] == "aux"


def test_bip_preview_oversize_upload_rejected_400(client, bip_dir, monkeypatch):
    """共享读取器超限异常映射 400 且沿用「上限」口径；不落预览文件。"""
    monkeypatch.setattr(bip_svc, "MAX_UPLOAD_BYTES", 16)
    _patch_reader(monkeypatch, bip_router)
    resp = _bip_preview(client, _xlsx_bytes(_aux_rows()), name="big.xlsx")
    assert resp.status_code == 400
    assert "上限" in resp.json()["detail"]


def test_tax_import_oversize_upload_rejected_400(client, monkeypatch):
    """税务导入超限同样 400 + 「上限」口径（与旧行为一致）。"""
    monkeypatch.setattr(tax_router, "MAX_UPLOAD_BYTES", 16)
    _patch_reader(monkeypatch, tax_router)
    resp = client.post("/api/finance/tax/financial-import/import",
                       files={"file": ("big.xlsx", b"x" * 64, XLSX_MIME)})
    assert resp.status_code == 400
    assert "上限" in resp.json()["detail"]


def test_upload_endpoints_fail_closed_when_shared_reader_missing(client, bip_dir, monkeypatch):
    """共享上传读取器缺失：明确 503，绝不退回无界 file.read()（A01 修复前行为）。"""
    monkeypatch.setattr(bip_router, "read_upload_limited", None, raising=False)
    resp = _bip_preview(client, _xlsx_bytes(_aux_rows()))
    assert resp.status_code == 503

    monkeypatch.setattr(tax_router, "read_upload_limited", None, raising=False)
    resp2 = client.post("/api/finance/tax/financial-import/import",
                        files={"file": ("r.xlsx", b"x" * 32, XLSX_MIME)})
    assert resp2.status_code == 503


# ---------------- 慢解析卸载：health 不被阻塞 ----------------

def test_bip_preview_parse_offloaded_event_loop_not_blocked(bip_dir, monkeypatch):
    """慢解析期间事件循环保持响应：BIP 预览的 Excel 解析经 asyncio.to_thread 卸载。"""
    real_read = bip_svc.read_bip_excel

    def slow_read(path):
        time.sleep(PARSE_SLEEP)
        return real_read(path)

    monkeypatch.setattr(bip_svc, "read_bip_excel", slow_read)
    _patch_reader(monkeypatch, bip_router)
    upload = UploadFile(filename="slow.xlsx", file=io.BytesIO(_xlsx_bytes(_aux_rows())))

    async def scenario():
        # 探针在解析可能占住循环之前就排入事件循环：同步解析时它必然被推迟到解析结束
        t0 = time.perf_counter()
        task = asyncio.create_task(bip_router.preview_import(upload))
        await asyncio.sleep(PROBE_SLEEP)
        probe = time.perf_counter() - t0
        result = await task
        return probe, result

    probe, result = asyncio.run(scenario())
    assert probe < PROBE_LIMIT, f"事件循环被同步解析阻塞：探针耗时 {probe:.3f}s"
    assert result["file_type"] == "aux" and result["row_count"] == 4


def test_tax_import_parse_offloaded_event_loop_not_blocked(monkeypatch):
    """慢解析期间事件循环保持响应：税务导入的解析经 asyncio.to_thread 卸载。"""

    def slow_import(file_path):
        time.sleep(PARSE_SLEEP)
        return tax_importer.ImportResult(success=True, summary="合成导入完成")

    monkeypatch.setattr(tax_router.importer, "import_from_file", slow_import)
    _patch_reader(monkeypatch, tax_router)
    upload = UploadFile(filename="slow.xlsx", file=io.BytesIO(_xlsx_bytes(_income_rows(2))))

    async def scenario():
        t0 = time.perf_counter()
        task = asyncio.create_task(tax_router.import_financial_statement(upload))
        await asyncio.sleep(PROBE_SLEEP)
        probe = time.perf_counter() - t0
        result = await task
        return probe, result

    probe, result = asyncio.run(scenario())
    assert probe < PROBE_LIMIT, f"事件循环被同步解析阻塞：探针耗时 {probe:.3f}s"
    assert result.success is True


# ---------------- 行数上限：超限无半成品 ----------------

def test_bip_read_bip_excel_enforces_row_limit(tmp_path, monkeypatch):
    """BIP 解析阶段可检测行数溢出：超限抛 ValueError；正常行数行为不变。"""
    path = tmp_path / "aux.xlsx"
    path.write_bytes(_xlsx_bytes(_aux_rows()))
    monkeypatch.setattr(bip_svc, "MAX_IMPORT_ROWS", 3, raising=False)
    with pytest.raises(ValueError, match="上限"):
        bip_svc.read_bip_excel(path)
    monkeypatch.setattr(bip_svc, "MAX_IMPORT_ROWS", 100, raising=False)
    assert len(bip_svc.read_bip_excel(path)) == 4


def test_bip_row_limit_rejects_confirm_without_partial_write(client, bip_dir, tmp_path, monkeypatch):
    """行数超限：确认导入 400 且不写任何批次/凭证；预览同样在解析阶段被拒。"""
    _patch_reader(monkeypatch, bip_router)
    payload = _xlsx_bytes(_aux_rows() * 2)  # 8 行
    token = _bip_preview(client, payload, name="many.xlsx").json()["token"]

    monkeypatch.setattr(bip_svc, "MAX_IMPORT_ROWS", 3, raising=False)
    resp = client.post("/api/finance/ledger/bip/imports/confirm", json={"token": token})
    assert resp.status_code == 400
    assert "上限" in resp.json()["detail"]
    # 无半成品：没有任何批次与凭证落库
    assert client.get("/api/finance/ledger/bip/imports").json()["total"] == 0
    assert client.get("/api/finance/ledger/bip/vouchers").json()["total"] == 0

    monkeypatch.setattr(bip_svc, "MAX_IMPORT_ROWS", 2, raising=False)
    resp2 = _bip_preview(client, payload, name="many.xlsx")
    assert resp2.status_code == 400
    assert "上限" in resp2.json()["detail"]


def test_tax_template_parser_detects_row_overflow(tmp_path, monkeypatch):
    """标准模板解析路径：单 sheet 行数超限即抛错（可检测溢出，不生成巨型文件）。"""
    path = tmp_path / "t.xlsx"
    path.write_bytes(_xlsx_bytes(_income_rows(6), sheet="利润表"))
    monkeypatch.setattr(tax_importer, "MAX_IMPORT_ROWS", 3, raising=False)
    limit_error = getattr(tax_importer, "ImportRowLimitExceeded", ValueError)
    with pytest.raises(limit_error, match="上限"):
        tax_importer.StandardTemplateParser.parse(str(path))


def test_tax_xml_parser_aborts_streaming_on_row_overflow(tmp_path, monkeypatch):
    """财务系统导出解析路径：流式逐行计数，超限即中止，不在内存累积全表。"""
    path = tmp_path / "x.xlsx"
    path.write_bytes(_xlsx_bytes(_income_rows(50), sheet="Sheet"))
    monkeypatch.setattr(tax_importer, "MAX_IMPORT_ROWS", 3, raising=False)
    limit_error = getattr(tax_importer, "ImportRowLimitExceeded", ValueError)
    with pytest.raises(limit_error, match="上限"):
        tax_importer.FinancialSystemXmlParser.parse(str(path))


def test_tax_import_row_limit_no_partial_result(client, monkeypatch):
    """行数超限：导入返回 success=False 与行数上限说明，不产出任何映射数据（无半成品）。"""
    _patch_reader(monkeypatch, tax_router)
    monkeypatch.setattr(tax_importer, "MAX_IMPORT_ROWS", 3, raising=False)
    resp = client.post("/api/finance/tax/financial-import/import",
                       files={"file": ("many.xlsx", _xlsx_bytes(_income_rows(6), sheet="利润表"), XLSX_MIME)})
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is False
    assert body["financial_data"] is None
    assert any("上限" in e["message"] for e in body["errors"])


# ---------------- subject_prefix LIKE 字面匹配 ----------------

def test_bip_subject_prefix_matches_wildcards_literally(client, bip_dir, tmp_path, monkeypatch):
    """subject_prefix 的 % / _ / 反斜杠按字面匹配，不通配放大结果集。"""
    _patch_reader(monkeypatch, bip_router)
    rows = _aux_rows() + [{
        "科目编码": r"10\0210", "科目名称": "银行存款-特别户", "日期": "2026-09-09",
        "核算账簿": "测试账簿甲", "凭证号": "记-0013", "摘要": "字面量测试",
        "借方本币": 10.0, "贷方本币": 0.0,
    }]
    token = _bip_preview(client, _xlsx_bytes(rows), name="lit.xlsx").json()["token"]
    assert client.post("/api/finance/ledger/bip/imports/confirm", json={"token": token}).status_code == 200

    base = "/api/finance/ledger/bip/entries"
    # 普通前缀语义不变
    assert client.get(base, params={"subject_prefix": "1002"}).json()["total"] == 2
    # 反斜杠按字面匹配：只命中编码以 10\02 开头的特别户（若转义不完整会误中 1002xx）
    assert client.get(base, params={"subject_prefix": r"10\02"}).json()["total"] == 1
    # % / _ 按字面匹配：数据中不存在这些前缀 → 0 条（修复前会匹配全部）
    assert client.get(base, params={"subject_prefix": "%"}).json()["total"] == 0
    assert client.get(base, params={"subject_prefix": "_"}).json()["total"] == 0


# ---------------- 税务导入错误不回带本机路径 ----------------

def test_tax_import_error_keeps_local_path_out_of_response(client, monkeypatch, caplog):
    """导入异常对外不回带本机异常路径；服务器保留完整异常日志。"""
    secret = r"C:\Users\someone\AppData\Local\Temp\tmp_secret_123.xlsx"

    def boom(file_path):
        raise RuntimeError(f"无法打开文件 {secret}")

    monkeypatch.setattr(tax_router.importer, "import_from_file", boom)
    _patch_reader(monkeypatch, tax_router)

    with caplog.at_level(logging.WARNING):
        resp = client.post("/api/finance/tax/financial-import/import",
                           files={"file": ("r.xlsx", _xlsx_bytes(_income_rows(2)), XLSX_MIME)})
    assert resp.status_code == 500
    assert "tmp_secret_123" not in resp.text
    assert secret not in resp.text
    assert "tmp_secret_123" in caplog.text  # 服务器日志保留完整异常


# ---------------- 二轮复核：对外异常脱敏与超限无残留 ----------------

def test_bip_preview_parse_error_keeps_local_path_out_of_response(client, bip_dir, monkeypatch, caplog):
    """BIP 预览解析异常对外不回带本机异常路径；服务器保留完整异常日志。"""
    secret = r"C:\Users\someone\AppData\Local\Temp\tmp_bip_preview_secret.xlsx"

    def boom(path):
        raise RuntimeError(f"无法打开文件 {secret}")

    monkeypatch.setattr(bip_svc, "read_bip_excel", boom)
    _patch_reader(monkeypatch, bip_router)

    with caplog.at_level(logging.WARNING):
        resp = _bip_preview(client, _xlsx_bytes(_aux_rows()))
    assert resp.status_code == 400
    assert secret not in resp.text
    assert "tmp_bip_preview_secret" not in resp.text
    assert "tmp_bip_preview_secret" in caplog.text  # 服务器日志保留完整异常


def test_bip_confirm_parse_error_keeps_local_path_out_of_response(client, bip_dir, monkeypatch, caplog):
    """BIP 确认解析异常同样脱敏（400 通用提示）；完整异常进服务器日志。"""
    secret = r"C:\Users\someone\AppData\Local\Temp\tmp_bip_confirm_secret.xlsx"
    _patch_reader(monkeypatch, bip_router)
    token = _bip_preview(client, _xlsx_bytes(_aux_rows()), name="ok.xlsx").json()["token"]

    def boom(path):
        raise RuntimeError(f"无法打开文件 {secret}")

    monkeypatch.setattr(bip_svc, "read_bip_excel", boom)
    with caplog.at_level(logging.WARNING):
        resp = client.post("/api/finance/ledger/bip/imports/confirm", json={"token": token})
    assert resp.status_code == 400
    assert secret not in resp.text and "tmp_bip_confirm_secret" not in resp.text
    assert "tmp_bip_confirm_secret" in caplog.text


def test_bip_preview_corrupt_file_rejected_generic_without_residue(client, bip_dir, monkeypatch):
    """损坏 Excel（非 zip 内容）：400 通用提示（不回带底层异常文本），tmp 不留半成品。"""
    _patch_reader(monkeypatch, bip_router)
    resp = _bip_preview(client, b"this is not an excel file at all", name="broken.xlsx")
    assert resp.status_code == 400
    detail = resp.json()["detail"]
    assert "Excel 解析失败" in detail
    assert "zip" not in detail.lower() and "BadZipFile" not in detail
    tmp_dir = bip_dir / "tmp"
    assert not (tmp_dir.exists() and any(tmp_dir.iterdir())), "解析失败不得遗留 tmp 预览文件"


def test_bip_row_limit_preview_leaves_no_tmp_residue(client, bip_dir, monkeypatch):
    """行数超限的预览在解析阶段被拒，且不在 tmp 留下被拒文件（无半成品）。"""
    _patch_reader(monkeypatch, bip_router)
    monkeypatch.setattr(bip_svc, "MAX_IMPORT_ROWS", 3, raising=False)
    resp = _bip_preview(client, _xlsx_bytes(_aux_rows() * 2), name="many.xlsx")  # 8 行
    assert resp.status_code == 400
    assert "上限" in resp.json()["detail"]
    tmp_dir = bip_dir / "tmp"
    assert not (tmp_dir.exists() and any(tmp_dir.iterdir())), "行数超限不得遗留 tmp 预览文件"


def test_bip_confirm_parse_failure_leaves_no_tmp_residue(client, bip_dir, monkeypatch):
    """确认入口解析失败（行数超限）：tmp 预览文件一并清理，无半成品残留。"""
    _patch_reader(monkeypatch, bip_router)
    token = _bip_preview(client, _xlsx_bytes(_aux_rows() * 2), name="many.xlsx").json()["token"]
    tmp_dir = bip_dir / "tmp"
    assert any(tmp_dir.iterdir())  # 预览成功时文件在（供确认入库使用）

    monkeypatch.setattr(bip_svc, "MAX_IMPORT_ROWS", 3, raising=False)
    resp = client.post("/api/finance/ledger/bip/imports/confirm", json={"token": token})
    assert resp.status_code == 400 and "上限" in resp.json()["detail"]
    assert not any(tmp_dir.iterdir()), "确认解析失败不得遗留 tmp 预览文件"


def test_tax_template_download_error_keeps_local_path_out_of_response(client, monkeypatch, caplog):
    """模板生成异常对外不回带本机异常路径（三个模板端点同口径）；服务器保留完整日志。"""
    secret = r"C:\Users\someone\AppData\Local\Temp\tmp_template_secret.xlsx"

    def boom(*args, **kwargs):
        raise RuntimeError(f"模板渲染失败：{secret}")

    monkeypatch.setattr(tax_router.importer, "generate_template", boom)
    for url in ("/api/finance/tax/financial-import/template",
                "/api/finance/tax/financial-import/template/balance-sheet",
                "/api/finance/tax/financial-import/template/income-statement"):
        with caplog.at_level(logging.WARNING):
            resp = client.get(url)
        assert resp.status_code == 500, url
        assert secret not in resp.text and "tmp_template_secret" not in resp.text
        assert "tmp_template_secret" in caplog.text  # 服务器日志保留完整异常


def test_tax_template_download_success_unchanged(client):
    """模板下载正常契约不变：200 + xlsx 三表（资产负债表含既有项目行）。"""
    resp = client.get("/api/finance/tax/financial-import/template")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in resp.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["资产负债表", "利润表", "现金流量表"]
    ws = wb["资产负债表"]
    items = [ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)]
    assert "货币资金" in items and "资产总计" in items and "所有者权益合计" in items


# ---------------- 共享接口接线（集成回归） ----------------

def test_shared_import_limits_wired_into_consumers():
    """消费侧确实接上共享上限（而非回退值/导入深度错误）：与 app.import_limits 同一对象。

    防的是相对导入层级写错（services/ 比模块根深一层）时被回退分支静默掩盖的问题。
    """
    from app.import_limits import MAX_IMPORT_ROWS

    assert bip_svc.MAX_IMPORT_ROWS is MAX_IMPORT_ROWS
    assert tax_importer.MAX_IMPORT_ROWS is MAX_IMPORT_ROWS


def test_shared_upload_reader_wired_into_routers():
    """两个上传路由确实接上共享读取器（而非 None 回退分支）。"""
    from app.upload import read_upload_limited as shared_reader

    assert bip_router.read_upload_limited is shared_reader
    assert tax_router.read_upload_limited is shared_reader


# ---------------- 正常导入行为不变 ----------------

def test_bip_preview_confirm_flow_unchanged(client, bip_dir, tmp_path, monkeypatch):
    """正常 预览→确认 全流程行为不变（共享读取 + to_thread 卸载不改变结果）。"""
    _patch_reader(monkeypatch, bip_router)
    token = _bip_preview(client, _xlsx_bytes(_aux_rows()), name="aux.xlsx").json()["token"]
    resp = client.post("/api/finance/ledger/bip/imports/confirm", json={"token": token})
    assert resp.status_code == 200, resp.text
    assert resp.json()["inserted_entries"] == 4
    entries = client.get("/api/finance/ledger/bip/entries", params={"subject_prefix": "1002"}).json()
    assert entries["total"] == 2


def test_tax_normal_import_unchanged(client, monkeypatch):
    """正常模板导入行为不变：200 + success + 财务数据映射。"""
    _patch_reader(monkeypatch, tax_router)
    payload = _xlsx_bytes([
        ["项目名称", "本期金额（万元）"],
        ["营业收入", 300.0], ["营业成本", 150.0], ["净利润", 100.0],
    ], sheet="利润表")
    resp = client.post("/api/finance/tax/financial-import/import",
                       files={"file": ("r.xlsx", payload, XLSX_MIME)})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["success"] is True
    assert body["format_detected"] == "standard_template"
    assert body["financial_data"]["total_revenue"] == 300.0
