"""AI 助手文档附件：正文提取、旧版转换、上传与模型消息链。"""
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from app import document_text as doc


def _make_pdf(path: Path, *page_texts: str) -> None:
    pdf = canvas.Canvas(str(path))
    for text in page_texts or ("PDF_DEFAULT_TEXT",):
        pdf.drawString(72, 720, text)
        pdf.showPage()
    pdf.save()


def _make_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("WORD_ALPHA_BEFORE_TABLE")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "WORD_BRAVO_CELL"
    table.cell(0, 1).text = "WORD_CHARLIE_CELL"
    document.add_paragraph("WORD_DELTA_AFTER_TABLE")
    document.save(path)


def _make_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "PPT_ALPHA_TITLE"
    slide.placeholders[1].text = "PPT_BRAVO_BODY"
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "PPT_CHARLIE_CELL"
    slide.notes_slide.notes_text_frame.text = "PPT_DELTA_NOTE"
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "PPT_ECHO_SECOND_SLIDE"
    presentation.save(path)


def test_pdf_extracts_page_text(tmp_path):
    source = tmp_path / "report.pdf"
    destination = tmp_path / "out" / "report.txt"
    _make_pdf(source, "PDF_ALPHA_PAGE_ONE", "PDF_BRAVO_PAGE_TWO")

    result = doc.extract_document(source, destination)

    text = destination.read_text(encoding="utf-8")
    assert "[第 1 页]" in text and "[第 2 页]" in text
    assert "PDF_ALPHA_PAGE_ONE" in text and "PDF_BRAVO_PAGE_TWO" in text
    assert result.characters == len(text)
    assert result.truncated is False and result.converted is False


def test_docx_extracts_paragraphs_and_tables_in_order(tmp_path):
    source = tmp_path / "report.docx"
    destination = tmp_path / "report.txt"
    _make_docx(source)

    result = doc.extract_document(source, destination)

    text = destination.read_text(encoding="utf-8")
    positions = [text.index(value) for value in (
        "WORD_ALPHA_BEFORE_TABLE", "WORD_BRAVO_CELL", "WORD_CHARLIE_CELL", "WORD_DELTA_AFTER_TABLE",
    )]
    assert positions == sorted(positions)
    assert "WORD_BRAVO_CELL\tWORD_CHARLIE_CELL" in text
    assert result.characters == len(text) and result.truncated is False


def test_pptx_extracts_slides_tables_and_notes(tmp_path):
    source = tmp_path / "deck.pptx"
    destination = tmp_path / "deck.txt"
    _make_pptx(source)

    result = doc.extract_document(source, destination)

    text = destination.read_text(encoding="utf-8")
    assert "幻灯片 1" in text and "幻灯片 2" in text
    for expected in (
        "PPT_ALPHA_TITLE", "PPT_BRAVO_BODY", "PPT_CHARLIE_CELL", "PPT_DELTA_NOTE", "PPT_ECHO_SECOND_SLIDE",
    ):
        assert expected in text
    assert result.characters == len(text) and result.truncated is False


@pytest.mark.parametrize("suffix", [".doc", ".ppt"])
def test_legacy_formats_reuse_ooxml_extraction(tmp_path, monkeypatch, suffix):
    modern = tmp_path / f"converted.{ 'docx' if suffix == '.doc' else 'pptx' }"
    if suffix == ".doc":
        _make_docx(modern)
    else:
        _make_pptx(modern)
    source = tmp_path / f"legacy{suffix}"
    source.write_bytes(b"legacy-binary")
    destination = tmp_path / "legacy.txt"

    def fake_convert(actual_source, work_dir):
        assert actual_source == source
        assert Path(work_dir).is_dir()
        return modern

    monkeypatch.setattr(doc, "_convert_legacy_to_ooxml", fake_convert)
    result = doc.extract_document(source, destination)

    assert result.converted is True
    assert ("WORD_ALPHA_BEFORE_TABLE" if suffix == ".doc" else "PPT_ALPHA_TITLE") in destination.read_text(encoding="utf-8")


def test_blank_pdf_raises_without_destination(tmp_path):
    source = tmp_path / "blank.pdf"
    destination = tmp_path / "blank.txt"
    _make_pdf(source, "")

    with pytest.raises(doc.DocumentExtractionError, match="未提取到文字"):
        doc.extract_document(source, destination)
    assert not destination.exists()


def test_malformed_docx_raises_without_destination(tmp_path):
    source = tmp_path / "fake.docx"
    destination = tmp_path / "fake.txt"
    source.write_bytes(b"not-an-office-file")

    with pytest.raises(doc.DocumentExtractionError, match="无法解析"):
        doc.extract_document(source, destination)
    assert not destination.exists()


def test_extraction_limits_and_truncation_marker(tmp_path, monkeypatch):
    source = tmp_path / "long.pdf"
    destination = tmp_path / "long.txt"
    _make_pdf(source, "PDF_PAGE_ONE", "PDF_PAGE_TWO", "PDF_PAGE_THREE")
    monkeypatch.setattr(doc, "MAX_DOCUMENT_TEXT_CHARS", 120)
    monkeypatch.setattr(doc, "MAX_PDF_PAGES", 2)

    result = doc.extract_document(source, destination)

    text = destination.read_text(encoding="utf-8")
    assert result.truncated is True
    assert result.characters == len(text) <= 120
    assert "已截断" in text and "PDF_PAGE_ONE" in text
    assert "PDF_PAGE_THREE" not in text


def test_ppt_slide_limit_is_explicitly_truncated(tmp_path, monkeypatch):
    source = tmp_path / "deck.pptx"
    destination = tmp_path / "deck.txt"
    _make_pptx(source)
    monkeypatch.setattr(doc, "MAX_PPT_SLIDES", 1)

    result = doc.extract_document(source, destination)

    text = destination.read_text(encoding="utf-8")
    assert result.truncated is True and "已截断" in text
    assert "PPT_ALPHA_TITLE" in text and "PPT_ECHO_SECOND_SLIDE" not in text


def test_ooxml_expansion_guard_rejects_before_parser(tmp_path, monkeypatch):
    source = tmp_path / "expanded.docx"
    destination = tmp_path / "expanded.txt"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("word/document.xml", b"x" * 2048)
    monkeypatch.setattr(doc, "MAX_OOXML_UNCOMPRESSED_BYTES", 1024)

    def parser_must_not_run(_source):
        raise AssertionError("OOXML 解析器在解压护栏前被调用")

    monkeypatch.setattr("docx.Document", parser_must_not_run)
    with pytest.raises(doc.DocumentExtractionError, match="解压"):
        doc.extract_document(source, destination)
    assert not destination.exists()


def test_legacy_converter_missing_is_clean_error(tmp_path, monkeypatch):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy")
    destination = tmp_path / "legacy.txt"
    monkeypatch.setattr(doc, "_find_soffice", lambda: None)

    with pytest.raises(doc.DocumentExtractionError, match="LibreOffice"):
        doc.extract_document(source, destination)
    assert not destination.exists()


def test_legacy_converter_timeout_is_clean_error(tmp_path, monkeypatch):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy")
    destination = tmp_path / "legacy.txt"
    monkeypatch.setattr(doc, "_find_soffice", lambda: "C:/mock/soffice.exe")

    def timeout(*_args, **_kwargs):
        raise subprocess.TimeoutExpired("soffice", 60)

    monkeypatch.setattr(doc.subprocess, "run", timeout)
    with pytest.raises(doc.DocumentExtractionError, match="超时"):
        doc.extract_document(source, destination)
    assert not destination.exists()


def test_legacy_converter_without_output_is_clean_error(tmp_path, monkeypatch):
    source = tmp_path / "legacy.ppt"
    source.write_bytes(b"legacy")
    destination = tmp_path / "legacy.txt"
    monkeypatch.setattr(doc, "_find_soffice", lambda: "C:/mock/soffice.exe")

    def no_output(*_args, **_kwargs):
        return subprocess.CompletedProcess("soffice", 0, stdout="", stderr="")

    monkeypatch.setattr(doc.subprocess, "run", no_output)
    with pytest.raises(doc.DocumentExtractionError, match="未生成"):
        doc.extract_document(source, destination)
    assert not destination.exists()


@pytest.fixture()
def upload_dir(tmp_path, monkeypatch):
    from app import config as app_config
    from app.routers import assistant as assistant_router

    target = tmp_path / "uploads"
    target.mkdir()
    monkeypatch.setattr(assistant_router, "UPLOAD_DIR", target)
    monkeypatch.setattr(app_config, "UPLOAD_DIR", target)
    return target


def _document_source(filename: str, tmp_path: Path) -> Path:
    source = tmp_path / "fixtures" / filename
    source.parent.mkdir(exist_ok=True)
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        _make_pdf(source)
    elif suffix == ".docx":
        _make_docx(source)
    elif suffix == ".pptx":
        _make_pptx(source)
    else:
        source.write_bytes(b"legacy-binary")
    return source


def _upload_document(client, source: Path, session_id="sess-doc-upload"):
    response = client.post(
        "/api/assistant/upload",
        data={"session_id": session_id},
        files={"file": (source.name, source.read_bytes())},
    )
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize("filename,expected", [
    ("report.pdf", "PDF_DEFAULT_TEXT"),
    ("report.docx", "WORD_ALPHA_BEFORE_TABLE"),
    ("report.pptx", "PPT_ALPHA_TITLE"),
])
def test_document_upload_preserves_original_and_writes_sidecar(
    client, upload_dir, tmp_path, filename, expected,
):
    source = _document_source(filename, tmp_path)

    body = _upload_document(client, source)

    original = Path(body["path"])
    assert body["name"] == filename and body["kind"] == "document_text"
    assert body["size"] == source.stat().st_size
    assert original == upload_dir / "sess-doc-upload" / filename
    assert original.is_file()
    assert "已提取" in body["note"]
    sidecar = doc.extracted_text_path(original)
    assert sidecar.is_file() and expected in sidecar.read_text(encoding="utf-8")


def test_user_content_blocks_contain_document_text_not_binary_fallback(client, upload_dir, tmp_path):
    from app.agents import orchestrator

    source = _document_source("report.pdf", tmp_path)
    uploaded = _upload_document(client, source)

    blocks = orchestrator._user_content_blocks(
        "请分析附件",
        [{"name": uploaded["name"], "path": uploaded["path"]}],
        None,
    )

    assert "PDF_DEFAULT_TEXT" in blocks[0].text
    assert "二进制或超大文件" not in blocks[0].text


def test_orchestrator_chat_passes_document_text_to_agent(client, upload_dir, tmp_path, monkeypatch):
    import asyncio
    from app.agents import orchestrator
    from app.agents.team import AgentSpec

    source = _document_source("report.pdf", tmp_path)
    uploaded = _upload_document(client, source)
    spec = AgentSpec(
        name="doc_reader", label="文档助手", emoji="", color="", description="", system_prompt="", model_id=None,
    )
    monkeypatch.setattr(orchestrator, "llm_status", lambda: {"configured": True})
    monkeypatch.setattr(orchestrator, "load_team", lambda: {"doc_reader": spec})
    monkeypatch.setattr(orchestrator, "_persist", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(orchestrator, "_touch_session", lambda *_args, **_kwargs: None)
    received = []

    class _Reply:
        def get_text_content(self):
            return "已分析"

        def get_content_blocks(self):
            return []

    class _Agent:
        async def reply(self, msg):
            received.append(msg)
            return _Reply()

    async def fake_ensure_agent(_session_id, _spec, _mode):
        return _Agent()

    monkeypatch.setattr(orchestrator, "_ensure_agent", fake_ensure_agent)
    result = asyncio.run(orchestrator.chat(
        session_id="sess-doc-chat",
        message="请分析附件",
        agent_name="doc_reader",
        attachments=[{"name": uploaded["name"], "path": uploaded["path"]}],
    ))

    assert result["ok"] is True and result["reply"] == "已分析"
    assert len(received) == 1
    assert "PDF_DEFAULT_TEXT" in received[0].get_content_blocks()[0].text


@pytest.mark.parametrize("kind,filename,error", [
    ("malformed", "fake.docx", "无法解析"),
    ("blank", "blank.pdf", "未提取到文字"),
    ("legacy", "legacy.doc", "LibreOffice"),
])
def test_document_upload_failure_is_400_and_cleans_original_and_sidecar(
    client, upload_dir, tmp_path, monkeypatch, kind, filename, error,
):
    if kind == "malformed":
        content = b"not-an-office-file"
    elif kind == "blank":
        blank = tmp_path / "blank-source.pdf"
        _make_pdf(blank, "")
        content = blank.read_bytes()
    else:
        content = b"legacy-binary"
        monkeypatch.setattr(doc, "_find_soffice", lambda: None)
    session_dir = upload_dir / "sess-doc-fail"
    sidecar = doc.extracted_text_path(session_dir / filename)
    sidecar.parent.mkdir(parents=True)
    sidecar.write_text("STALE_EXTRACTION", encoding="utf-8")

    response = client.post(
        "/api/assistant/upload",
        data={"session_id": "sess-doc-fail"},
        files={"file": (filename, content)},
    )

    assert response.status_code == 400, response.text
    assert error in response.json()["detail"]
    assert not any(path.is_file() for path in session_dir.rglob("*"))


def test_document_sidecar_cannot_escape_upload_root(upload_dir, tmp_path):
    from app.agents import orchestrator

    outside = tmp_path / "outside.pdf"
    _make_pdf(outside, "OUTSIDE_SECRET_TEXT")
    sidecar = doc.extracted_text_path(outside)
    sidecar.parent.mkdir()
    sidecar.write_text("OUTSIDE_SECRET_TEXT", encoding="utf-8")

    block = orchestrator._attachments_block([{"name": "outside.pdf", "path": str(outside)}])

    assert "OUTSIDE_SECRET_TEXT" not in block


def test_document_without_sidecar_asks_for_reupload(client, upload_dir, tmp_path):
    from app.agents import orchestrator

    original = upload_dir / "sess-doc-missing" / "report.pdf"
    original.parent.mkdir()
    _make_pdf(original)

    block = orchestrator._attachments_block([{"name": "report.pdf", "path": str(original)}])

    assert "提取文本不存在，请重新上传" in block
    assert "二进制或超大文件" not in block
